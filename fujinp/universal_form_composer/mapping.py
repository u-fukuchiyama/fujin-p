"""ゆにこん — 部品と常設 SQL テーブルの対応付け。

対応付けはプロジェクトの JSON の mapping に，SQL テーブルと同じ形の格子として持つ。
行は SQL の行（キーの組），セルは中身の代わりにこんかの部品 ID（または key_path）。

  details … plan_details と同型。行＝(年度, 中期計画番号, 年度計画番号, 細目番号)。
            セルは部品 ID の並び。本文の k 行目（「・」1つ）がセルの k 番目の部品に対応する。
            同じ部品を複数の細目のセルに書いてよい（一つの「・」を複数のグループが受け持つ場合）。
  plans   … plan_evaluations と同型。行＝(年度, 中期計画番号, 年度計画番号)。セルは部品 ID 1つ。

担当（owner_group）は ID ではなく値（まいぐるのグループ名）として格子に書く。
"""

import io
import re

from . import store

DETAIL_KEYS = [('fiscal_year', '年度'), ('mid_plan_no', '中期計画番号'),
               ('annual_plan_no', '年度計画番号'), ('detail_no', '細目番号')]
DETAIL_VALS = [('owner_group', '担当')]
DETAIL_CELLS = [('plan_text', '細目計画内容'), ('progress_text', '細目進捗内容'),
                ('result_text', '細目業務実績内容'), ('self_eval', '自己評価'),
                ('self_eval_reason', '自己評価の理由'), ('editor_memo', '編集者メモ'),
                ('owner_memo', '担当者メモ')]
PLAN_KEYS = DETAIL_KEYS[:3]
PLAN_CELLS = [('self_eval', '自己評価'), ('self_eval_reason', '自己評価の理由'),
              ('committee_eval', '評価委員会評価'), ('committee_comment', '評価委員会質問とコメント'),
              ('response_to_committee', '評価委員会への回答'), ('editor_memo', '編集者メモ')]

MID_DETAIL_KEYS = [('term', '期'), ('sub_term', '版'), ('mid_plan_no', '中期計画番号'), ('detail_no', '細目番号')]
MID_DETAIL_VALS = [('mid_goal_no', '中期目標番号'), ('owner_group', '担当')]
MID_DETAIL_CELLS = [('mid_goal_text', '中期目標'), ('mid_plan_text', '中期計画（細目）'),
                    ('progress_text', '細目進捗内容'), ('result_text', '細目業務実績内容'),
                    ('self_eval', '自己評価'), ('self_eval_reason', '自己評価の理由'),
                    ('editor_memo', '編集者メモ'), ('owner_memo', '担当者メモ')]
MID_EVAL_KEYS = [('term', '期'), ('sub_term', '版'), ('mid_plan_no', '中期計画番号'), ('phase', '時期')]
MID_EVAL_CELLS = PLAN_CELLS

KINDS = ('details', 'plans', 'mid_details', 'mid_evals')
SHEETS = {'details': '細目', 'plans': '計画番号', 'mid_details': '中期細目', 'mid_evals': '中期評価'}
TABS = {'details': '年度計画の細目（T_ann_plan_details）', 'plans': '年度計画番号（T_ann_plan_evaluations）',
        'mid_details': '中期計画の細目（T_mid_plan_details）', 'mid_evals': '中期計画の評価（T_mid_plan_evaluations）'}
STR_KEYS = {'phase': ('progress', 'result')}     # 整数でないキー（取りうる値）
INT_VALS = {'mid_goal_no'}                        # 値の列のうち整数のもの


def spec(kind):
    if kind == 'details':
        return DETAIL_KEYS, DETAIL_VALS, DETAIL_CELLS
    if kind == 'plans':
        return PLAN_KEYS, [], PLAN_CELLS
    if kind == 'mid_details':
        return MID_DETAIL_KEYS, MID_DETAIL_VALS, MID_DETAIL_CELLS
    return MID_EVAL_KEYS, [], MID_EVAL_CELLS


def key_ok(v, col):
    """キーの値がそろっているか。"""
    if col in STR_KEYS:
        return v in STR_KEYS[col]
    return isinstance(v, int)


# ------------------------------------------------------------------ 値の読み方

def parse_ids(v):
    """'1203, 1204' / '#1203、#1204' / 改行区切り / 数値 → ['1203', '1204']。key_path もそのまま通す。"""
    if v is None:
        return []
    if isinstance(v, (int, float)):
        return [str(int(v))]
    if isinstance(v, (list, tuple)):
        out = []
        for x in v:
            out.extend(parse_ids(x))
        return out
    out = []
    for tok in re.split(r'[,\s、，;；]+', str(v)):
        tok = tok.strip().lstrip('#＃')
        if not tok:
            continue
        if re.fullmatch(r'\d+(\.0+)?', tok):
            tok = str(int(float(tok)))
        out.append(tok)
    return out


def _int(v):
    if v is None or v == '':
        return None
    try:
        return int(float(str(v).strip()))
    except (TypeError, ValueError):
        return 'bad:%s' % v


def empty():
    m = {'version': 1, 'updated_at': '', 'updated_by': None}
    for k in KINDS:
        m[k] = []
    return m


def get(data):
    m = (data or {}).get('mapping')
    if not isinstance(m, dict) or 'details' not in m:
        m = empty()
    for k in KINDS:
        m.setdefault(k, [])
    return m


def normalize_rows(kind, rows):
    """画面や xlsx から来た行をそろえる。空行は捨てる。"""
    keys, vals, cells = spec(kind)
    out = []
    for r in rows or []:
        if not isinstance(r, dict):
            continue
        row = {}
        for k, _ in keys:
            if k in STR_KEYS:
                v = str(r.get(k) or '').strip().lower()
                row[k] = v if v in STR_KEYS[k] else (('bad:%s' % v) if v else None)
            else:
                row[k] = _int(r.get(k))
        for k, _ in vals:
            if k in INT_VALS:
                row[k] = _int(r.get(k))
            else:
                row[k] = (str(r.get(k)).strip() if r.get(k) not in (None, '') else '')
        cs = r.get('cells') if isinstance(r.get('cells'), dict) else r
        row['cells'] = {k: parse_ids(cs.get(k)) for k, _ in cells}
        if all(row[k] is None for k, _ in keys) and not any(row['cells'].values()) \
                and not any(row.get(k) for k, _ in vals):
            continue
        out.append(row)

    def sk(r):
        return tuple(((0, r[k]) if isinstance(r[k], int) else
                      (1, STR_KEYS[k].index(r[k])) if (k in STR_KEYS and r[k] in STR_KEYS[k]) else (2, 0))
                     for k, _ in keys)
    out.sort(key=sk)
    return out


# ------------------------------------------------------------------ 点検

def brief(pj, ref):
    """部品の短い説明（画面の手掛かり・点検の文面用）。"""
    q = pj.resolve(ref)
    if not q:
        return None
    lv = pj.level_of(q)
    body = q.get('body') or ''
    if lv == 'box' and q.get('recipe') != 'text':
        kids = pj.children(q['id'])
        body = ' / '.join((k.get('body') or '') for k in kids[:3])
    snip = re.sub(r'\s+', ' ', re.sub(r'<[^>]+>', '', body)).strip()[:50]
    return {'id': q['id'], 'level': lv, 'key_path': q.get('key_path') or '',
            'title': q.get('title') or '', 'snippet': snip}


def check(pj, m, group_names=None):
    """対応付けの点検。文面の並び（errors, warnings, infos）と，セルごとの手掛かりを返す。"""
    errors, warnings, infos = [], [], []
    used = {}          # 部品 ID → [(種類, 行の名前, 列)]
    tips = {k: [] for k in KINDS}
    for kind in KINDS:
        keys, vals, cells = spec(kind)
        seen = {}
        for i, r in enumerate(m.get(kind) or []):
            kv = [r.get(k) for k, _ in keys]
            name = '%s %s' % (SHEETS[kind], '-'.join(
                '?' if v is None else (v[4:] if isinstance(v, str) and v.startswith('bad:') else str(v)) for v in kv))
            for (k, lb), v in zip(keys, kv):
                if v is None:
                    errors.append('%s：%sが空です' % (name, lb))
                elif k in STR_KEYS and not key_ok(v, k):
                    errors.append('%s：%sは %s のどれかです（%s）' % (name, lb, '／'.join(STR_KEYS[k]), str(v)[4:]))
                elif k not in STR_KEYS and isinstance(v, str):
                    errors.append('%s：%sが整数ではありません（%s）' % (name, lb, v[4:]))
            if 'detail_no' in dict(keys) and isinstance(r.get('detail_no'), int) and r['detail_no'] < 1:
                errors.append('%s：細目番号は1から始めます' % name)
            for k, lb in vals:
                if k in INT_VALS and not isinstance(r.get(k), int):
                    errors.append('%s：%sが整数で入っていません' % (name, lb))
            key = tuple(kv)
            if key in seen:
                errors.append('%s：同じキーの行が重複しています' % name)
            seen[key] = True
            if 'owner_group' in dict(vals) and group_names is not None:
                g = r.get('owner_group') or ''
                if not g:
                    warnings.append('%s：担当が空です' % name)
                elif g not in group_names:
                    warnings.append('%s：担当「%s」はまいぐるのグループにありません' % (name, g))
            rt = {}
            for k, lb in cells:
                ids = (r.get('cells') or {}).get(k) or []
                cell_tips, bad = [], False
                if kind in ('plans', 'mid_evals') and len(ids) > 1:
                    errors.append('%s・%s：計画番号の欄に書ける部品は1つです（%d個）' % (name, lb, len(ids)))
                    bad = True
                for ref in ids:
                    b = brief(pj, ref)
                    if not b:
                        errors.append('%s・%s：部品 %s が見つかりません' % (name, lb, ref))
                        cell_tips.append('%s：見つかりません' % ref)
                        bad = True
                        continue
                    if b['level'] in ('work', 'part'):
                        warnings.append('%s・%s：#%s は%sで，中身を入れる部品ではありません'
                                        % (name, lb, b['id'], '作品' if b['level'] == 'work' else '複合部品'))
                        bad = True
                    used.setdefault(b['id'], []).append((kind, name, k))
                    cell_tips.append('#%s %s %s' % (b['id'], b['key_path'] or b['title'], b['snippet']))
                rt[k] = {'tip': '\n'.join(cell_tips), 'bad': bad}
            tips[kind].append(rt)
    for qid, places in used.items():
        cols = {(kd, c) for kd, _, c in places}
        if len(cols) > 1:
            warnings.append('#%s が異なる欄に書かれています（%s）' % (
                qid, '，'.join(sorted({'%s・%s' % (SHEETS[kd], c) for kd, c in cols}))))
        elif len(places) > 1:
            infos.append('#%s は %d 行で共有されています（%s）' % (qid, len(places), '，'.join(p[1] for p in places[:4])
                                                               + ('…' if len(places) > 4 else '')))
    dkeys = {(r.get('fiscal_year'), r.get('mid_plan_no'), r.get('annual_plan_no')) for r in m.get('details') or []}
    pkeys = {(r.get('fiscal_year'), r.get('mid_plan_no'), r.get('annual_plan_no')) for r in m.get('plans') or []}
    for k in sorted(dkeys - pkeys, key=lambda x: tuple(v if isinstance(v, int) else 10 ** 9 for v in x)):
        if all(isinstance(v, int) for v in k):
            infos.append('計画番号 %s の行がありません（細目の行はあります）' % '-'.join(str(v) for v in k))
    return {'errors': errors, 'warnings': warnings, 'infos': infos, 'tips': tips,
            'n_details': len(m.get('details') or []), 'n_plans': len(m.get('plans') or []),
            'n_mid_details': len(m.get('mid_details') or []), 'n_mid_evals': len(m.get('mid_evals') or []),
            'n_parts': len(used)}


# ------------------------------------------------------------------ xlsx

def _label(col, lb):
    return '%s（%s）' % (lb, col)


def _col_of(header, kind):
    """見出しから列名を引く。「年度（fiscal_year）」「fiscal_year」「年度」のどれでもよい。"""
    keys, vals, cells = spec(kind)
    h = str(header or '').strip()
    m = re.search(r'[（(]\s*([a-z_]+)\s*[)）]', h)
    if m:
        return m.group(1)
    for k, lb in keys + vals + cells:
        if h in (k, lb):
            return k
    return None


def to_xlsx(pj, m):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws0 = wb.active
    ws0.title = '説明'
    notes = [
        'ゆにこん 対応付けシート',
        '「細目」「計画番号」「中期細目」「中期評価」の4シートが，常設テーブル T_ann_plan_details・T_ann_plan_evaluations・T_mid_plan_details・T_mid_plan_evaluations と同じ形の格子です．',
        'キーの列（年度・期・版・中期計画番号・年度計画番号・細目番号・時期）と，担当・中期目標番号には値を書きます．担当はまいぐるのグループ名です．時期は progress（4年終了時）か result（全期間終了時）です．',
        'そのほかの列には中身ではなく，こんかの部品 ID（または key_path）を書きます．',
        '細目の欄には部品 ID を複数，順に書けます（例：1203, 1204）．本文の k 行目の「・」が k 番目の部品に対応します．',
        '同じ部品を複数の細目の行に書いてかまいません（一つの「・」を複数のグループが受け持つとき）．',
        '計画番号・中期評価の欄に書ける部品は1つです．',
        '細目番号は1から始めます．見出し行（1行目）は変えないでください．',
        '「部品」シートは，このプロジェクトの部品の一覧です（手掛かり用．取り込みでは読みません）．',
    ]
    for i, t in enumerate(notes, 1):
        ws0.cell(row=i, column=1, value=t)
    ws0['A1'].font = Font(bold=True, size=13)
    ws0.column_dimensions['A'].width = 110
    head_fill = PatternFill('solid', fgColor='EFEAF9')
    key_fill = PatternFill('solid', fgColor='F6F5FA')
    for kind in KINDS:
        keys, vals, cells = spec(kind)
        cols = keys + vals + cells
        ws = wb.create_sheet(SHEETS[kind])
        for j, (k, lb) in enumerate(cols, 1):
            c = ws.cell(row=1, column=j, value=_label(k, lb))
            c.font = Font(bold=True); c.fill = head_fill
            c.alignment = Alignment(wrap_text=True, vertical='top')
            ws.column_dimensions[c.column_letter].width = 10 if (k, lb) in keys else (16 if (k, lb) in vals else 18)
        for i, r in enumerate(m.get(kind) or [], 2):
            j = 1
            for k, _ in keys:
                v = r.get(k)
                if isinstance(v, str) and v.startswith('bad:'):
                    v = v[4:]
                c = ws.cell(row=i, column=j, value=v)
                c.fill = key_fill; j += 1
            for k, _ in vals:
                v = r.get(k)
                if isinstance(v, str) and v.startswith('bad:'):
                    v = v[4:]
                ws.cell(row=i, column=j, value=v if v not in ('', None) else None); j += 1
            for k, _ in cells:
                ids = (r.get('cells') or {}).get(k) or []
                ws.cell(row=i, column=j, value=', '.join(ids) if ids else None); j += 1
        ws.freeze_panes = ws.cell(row=2, column=len(keys) + 1)
    wp = wb.create_sheet('部品')
    head = ['ID', '種別', 'key_path', '題名', '作法', '親（箇条のとき）', '書き出し']
    for j, h in enumerate(head, 1):
        c = wp.cell(row=1, column=j, value=h); c.font = Font(bold=True); c.fill = head_fill
    parent_of = {}
    for pid, lst in pj.kids.items():
        for ln in lst:
            parent_of.setdefault(int(ln['child']), pid)
    order = sorted(pj.q.values(), key=lambda q: q['id'])
    lvname = {'work': '作品', 'part': '複合部品', 'box': '倉庫部品', 'item': '箇条'}
    i = 2
    for q in order:
        b = brief(pj, q['id'])
        wp.cell(row=i, column=1, value=q['id'])
        wp.cell(row=i, column=2, value=lvname.get(b['level'], ''))
        wp.cell(row=i, column=3, value=b['key_path'] or None)
        wp.cell(row=i, column=4, value=b['title'] or None)
        wp.cell(row=i, column=5, value=q.get('recipe') or None)
        wp.cell(row=i, column=6, value=parent_of.get(q['id']))
        wp.cell(row=i, column=7, value=b['snippet'] or None)
        i += 1
    for col, w in zip('ABCDEFG', (8, 10, 26, 30, 8, 10, 60)):
        wp.column_dimensions[col].width = w
    wp.freeze_panes = 'A2'
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def from_xlsx(raw):
    """xlsx から対応付けを読む。(mapping, エラー文)"""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as e:
        return None, 'xlsx として読めません：%s' % e
    m = empty()
    m['_sheets'] = []
    for kind, sheet in SHEETS.items():
        if sheet not in wb.sheetnames:
            continue
        m['_sheets'].append(kind)
        ws = wb[sheet]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        cols = [_col_of(h, kind) for h in rows[0]]
        if not any(cols):
            return None, '「%s」シートの見出し行が読めません' % sheet
        out = []
        for r in rows[1:]:
            d = {}
            for c, v in zip(cols, r):
                if c:
                    d[c] = v
            out.append(d)
        m[kind] = normalize_rows(kind, out)
    if not m['_sheets']:
        return None, '「細目」「計画番号」「中期細目」「中期評価」のどのシートも見つかりません'
    return m, None


def all_group_names():
    try:
        with store._db('default') as (cur, conn):
            cur.execute('SELECT name FROM user_groups')
            return {r['name'] for r in cur.fetchall() if r.get('name')}
    except Exception:
        return None
