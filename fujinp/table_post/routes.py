# SPDX-FileCopyrightText: 2024-2026 Toyoaki Nishida
# SPDX-License-Identifier: AGPL-3.0-or-later
#
# This file is part of FUJIN-P.
# Copyright (C) 2024-2026 Toyoaki Nishida
#
# FUJIN-P is free software: you can redistribute it and/or modify
# it under the terms of the GNU Affero General Public License as published by
# the Free Software Foundation, either version 3 of the License, or
# (at your option) any later version.
#
# FUJIN-P is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
# GNU Affero General Public License for more details.
#
# You should have received a copy of the GNU Affero General Public License
# along with FUJIN-P.  If not, see <https://www.gnu.org/licenses/>.
#
# Source: https://github.com/u-fukuchiyama/fujin-p

# table_post/routes.py
# 【FUJIN-P テーブルポスト - 新プラットフォーム対応版】
#
# テーブルサイクルを二段階・公開運用向けに再設計したもの。
#   - 一般ユーザ : テーブルのxlsxダウンロード / 編集済xlsxの「更新申請」
#   - admin      : プロジェクト管理 / 申請の適用(本番更新) / ロールバック
#
# 新設テーブル(default DB):
#   table_post_projects / table_post_status / table_post_requests / table_post_history

import io
import json
import datetime
import logging
from pytz import timezone
import mysql.connector
import pandas as pd
import numpy as np
from flask import (Blueprint, request, jsonify, send_file, session,
                   render_template)
from decorators import login_required
from auth import redirect_to_dashboard
from config import Config
from db import DatabaseConfig
from notifiers import notify_channel   # FUJIN-P 共通通知サブシステム（トップレベル）
from . import table_post_bp

logging.basicConfig(level=logging.DEBUG)

# 新設テーブル名（default DBに配置）
T_PROJECTS = "table_post_projects"
T_STATUS   = "table_post_status"
T_REQUESTS = "table_post_requests"
T_HISTORY  = "table_post_history"


# ============================================================
# 共通ヘルパ
# ============================================================
JST = timezone('Asia/Tokyo')


def get_jst_now():
    """
    現在のJST日時（naive datetime）。FUJIN-P標準の日時方式。
    DBへ書く日時は必ずこれを使う（MySQLのNOW()やサーバローカル時刻は
    PythonAnywhereではUTCになるため使わない）。
    """
    return datetime.datetime.now(JST).replace(tzinfo=None)


def _is_admin():
    return session.get('user_category') in ('admin',)


def _current_user():
    """更新者として記録するアカウント名"""
    return (session.get('username')
            or session.get('user_name')
            or session.get('user_id')
            or 'unknown')


def _to_python(v):
    """numpy/pandas型 → Python標準型。NaN/NaT → None。"""
    try:
        if v is None:
            return None
        if not isinstance(v, str) and pd.isna(v):
            return None
    except (ValueError, TypeError):
        pass
    if isinstance(v, np.integer):
        return int(v)
    if isinstance(v, np.floating):
        return float(v)
    if isinstance(v, np.bool_):
        return bool(v)
    if hasattr(v, 'item'):
        return v.item()
    if isinstance(v, (datetime.datetime, datetime.date, datetime.time,
                      pd.Timestamp, datetime.timedelta, pd.Timedelta)):
        return str(v)
    return v


def _get_table_schema(cursor, table_name):
    """テーブルのカラム名リスト（定義順）を返す"""
    cursor.execute(f"DESCRIBE `{table_name}`")
    return [row[0] for row in cursor.fetchall()]


def _get_table_constraints(conn, table_name):
    """
    対象テーブルの制約情報を返す。
      col_info: {列名: {'type','nullable','default','auto'}}
      uniques : {索引名: [列名, ...]}   PRIMARY と UNIQUE 索引（複合含む）
    """
    cursor = conn.cursor(dictionary=True)
    try:
        cursor.execute(f"SHOW COLUMNS FROM `{table_name}`")
        col_info = {}
        for r in cursor.fetchall():
            col_info[r['Field']] = {
                'type': (r['Type'] or '').lower(),
                'nullable': r['Null'] == 'YES',
                'default': r['Default'],
                'auto': 'auto_increment' in (r['Extra'] or '').lower(),
            }
        cursor.execute(f"SHOW INDEX FROM `{table_name}` WHERE Non_unique = 0")
        idx = {}
        for r in cursor.fetchall():
            idx.setdefault(r['Key_name'], []).append(
                (r['Seq_in_index'], r['Column_name']))
        uniques = {k: [c for _, c in sorted(v)] for k, v in idx.items()}
        return col_info, uniques
    finally:
        cursor.close()


_STRING_TYPES = ('char', 'varchar', 'text', 'tinytext', 'mediumtext',
                 'longtext', 'enum', 'set')


def _fmt_val(v):
    """エラーメッセージ表示用（2.0 → 2 など）"""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return repr(v) if isinstance(v, str) else str(v)


def _validate_rows(columns, rows, col_info, uniques, max_show=10):
    """
    行データを対象テーブルの制約で点検する。
    戻り値: (rows, filled, problems)
      rows     : 補正後の行（NOT NULL 列の空値を '' または既定値に置換）
      filled   : 補正したセル数
      problems : 違反の説明文リスト（空なら合格）
    補正の理由: xlsx へ書き出した '' は読み戻すと NaN→None になるため、
    NOT NULL の文字列列では '' に戻す（往復で壊れないようにする）。
    非文字列列は既定値があればそれを使い、無ければ違反として報告する。
    Excel の行番号はヘッダを1行目として index+2 で表示する。
    """
    problems = []
    filled = 0
    null_hits = {}   # 列名 → [Excel行番号]

    fixed = []
    for i, row in enumerate(rows):
        row = list(row)
        for j, col in enumerate(columns):
            if row[j] is not None:
                continue
            info = col_info.get(col)
            if info is None or info['nullable'] or info['auto']:
                continue
            if info['type'].startswith(_STRING_TYPES):
                row[j] = ''
                filled += 1
            elif info['default'] is not None:
                row[j] = info['default']
                filled += 1
            else:
                null_hits.setdefault(col, []).append(i + 2)
        fixed.append(row)

    for col, lines in null_hits.items():
        head = ', '.join(str(x) for x in lines[:max_show])
        more = f' …他 {len(lines) - max_show} 行' if len(lines) > max_show else ''
        problems.append(f'[NOT NULL] 列「{col}」が空: {len(lines)} 行'
                        f'（行 {head}{more}）')

    col_pos = {c: k for k, c in enumerate(columns)}
    for key_name, key_cols in uniques.items():
        if not all(c in col_pos for c in key_cols):
            continue
        seen = {}
        for i, row in enumerate(fixed):
            key = tuple(row[col_pos[c]] for c in key_cols)
            seen.setdefault(key, []).append(i + 2)
        dups = {k: v for k, v in seen.items() if len(v) > 1}
        if not dups:
            continue
        total = sum(len(v) for v in dups.values())
        lines = [f'[{key_name}] キー({", ".join(key_cols)}) の重複: '
                 f'{len(dups)} 組 / {total} 行']
        for n, (k, v) in enumerate(dups.items()):
            if n >= max_show:
                lines.append(f'  …他 {len(dups) - max_show} 組')
                break
            kv = ', '.join(_fmt_val(x) for x in k)
            rv = ', '.join(str(x) for x in v[:8])
            if len(v) > 8:
                rv += f' …計 {len(v)} 行'
            lines.append(f'  ({kv}) : 行 {rv}')
        problems.append('\n'.join(lines))

    return fixed, filled, problems


def _check_payload(database, table_name, columns, rows):
    """制約点検の入口。DB接続して制約を取り、(rows, filled, problems) を返す。"""
    conn = mysql.connector.connect(**DatabaseConfig.get_config(database))
    try:
        col_info, uniques = _get_table_constraints(conn, table_name)
    finally:
        if conn.is_connected():
            conn.close()
    return _validate_rows(columns, rows, col_info, uniques)


def _read_table_as_dict(database, table_name):
    """テーブル全内容を {'columns': [...], 'rows': [[...], ...]} で返す"""
    conn = mysql.connector.connect(**DatabaseConfig.get_config(database))
    try:
        cursor = conn.cursor()
        columns = _get_table_schema(cursor, table_name)
        cursor.execute(f"SELECT * FROM `{table_name}`")
        rows = [[_to_python(v) for v in row] for row in cursor.fetchall()]
        return {'columns': columns, 'rows': rows}
    finally:
        if conn.is_connected():
            cursor.close()
            conn.close()


def _restore_table(database, table_name, columns, rows):
    """指定内容でテーブルを完全置換（DELETE → INSERT、外部キー制約は一時無効化）"""
    conn = mysql.connector.connect(**DatabaseConfig.get_config(database))
    try:
        cursor = conn.cursor()
        cursor.execute("SET FOREIGN_KEY_CHECKS = 0")
        conn.start_transaction()
        cursor.execute(f"DELETE FROM `{table_name}`")
        if rows:
            placeholders = ', '.join(['%s'] * len(columns))
            col_clause = '`, `'.join(str(c) for c in columns)
            insert_sql = (f"INSERT INTO `{table_name}` "
                          f"(`{col_clause}`) VALUES ({placeholders})")
            values = [tuple(_to_python(v) for v in row) for row in rows]
            cursor.executemany(insert_sql, values)
        cursor.execute("SET FOREIGN_KEY_CHECKS = 1")
        conn.commit()
        return len(rows)
    except Exception:
        conn.rollback()
        try:
            cursor.execute("SET FOREIGN_KEY_CHECKS = 1")
        except Exception:
            pass
        raise
    finally:
        if conn.is_connected():
            cursor.close()
            conn.close()


# ============================================================
# 画面
# ============================================================
@table_post_bp.route('/')
@login_required
def table_post_dashboard():
    """テーブルポスト メイン画面（一般ユーザ/admin共通、admin機能はJSで出し分け）"""
    return render_template('table_post/table_post_dashboard.html',
                           is_admin=_is_admin())


# ============================================================
# プロジェクト
# ============================================================
@table_post_bp.route('/get_projects', methods=['GET'])
@login_required
def get_projects():
    """
    プロジェクト一覧。
    通常は有効なもののみ。?include_inactive=1 で無効も含める（admin編集用）。
    """
    include_inactive = request.args.get('include_inactive') == '1'
    try:
        conn = mysql.connector.connect(**DatabaseConfig.default())
        cursor = conn.cursor(dictionary=True)
        sql = (f"SELECT id, project_name, description, created_by, "
               f"created_at, is_active FROM {T_PROJECTS}")
        if not include_inactive:
            sql += " WHERE is_active = 1"
        sql += " ORDER BY id DESC"
        cursor.execute(sql)
        projects = cursor.fetchall()
        return jsonify({'success': True, 'projects': projects})
    except mysql.connector.Error as e:
        logging.error("get_projects error: %s", e)
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        if 'conn' in locals() and conn.is_connected():
            cursor.close()
            conn.close()


@table_post_bp.route('/create_project', methods=['POST'])
@login_required
def create_project():
    """プロジェクト作成（admin限定）"""
    if not _is_admin():
        return jsonify({'success': False, 'error': '管理者権限が必要です'}), 403

    data = request.json or {}
    name = (data.get('project_name') or '').strip()
    description = (data.get('description') or '').strip()
    if not name:
        return jsonify({'success': False, 'error': 'プロジェクト名が未指定です'}), 400

    try:
        conn = mysql.connector.connect(**DatabaseConfig.default())
        cursor = conn.cursor()
        now = get_jst_now().strftime('%Y-%m-%d %H:%M:%S')
        cursor.execute(
            f"INSERT INTO {T_PROJECTS} "
            f"(project_name, description, created_by, created_at, is_active) "
            f"VALUES (%s, %s, %s, %s, 1)",
            (name, description, _current_user(), now))
        conn.commit()
        return jsonify({'success': True, 'project_id': cursor.lastrowid})
    except mysql.connector.Error as e:
        logging.error("create_project error: %s", e)
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        if 'conn' in locals() and conn.is_connected():
            cursor.close()
            conn.close()


@table_post_bp.route('/update_project', methods=['POST'])
@login_required
def update_project():
    """
    プロジェクトの編集（admin限定）。
    名称・説明の変更、および有効/無効(is_active)の切り替えを行う。
    無効化してもプロジェクト・状態・申請・記録の各データは削除されない
    （get_projects で一覧に出なくなるだけ）。
    """
    if not _is_admin():
        return jsonify({'success': False, 'error': '管理者権限が必要です'}), 403

    data = request.json or {}
    project_id = data.get('project_id')
    name = (data.get('project_name') or '').strip()
    description = (data.get('description') or '').strip()
    is_active = data.get('is_active')
    if not project_id:
        return jsonify({'success': False,
                        'error': 'project_id が未指定です'}), 400
    if not name:
        return jsonify({'success': False,
                        'error': 'プロジェクト名が未指定です'}), 400

    try:
        conn = mysql.connector.connect(**DatabaseConfig.default())
        cursor = conn.cursor()
        cursor.execute(
            f"UPDATE {T_PROJECTS} "
            f"SET project_name = %s, description = %s, is_active = %s "
            f"WHERE id = %s",
            (name, description, 1 if is_active else 0, project_id))
        conn.commit()
        return jsonify({'success': True})
    except mysql.connector.Error as e:
        logging.error("update_project error: %s", e)
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        if 'conn' in locals() and conn.is_connected():
            cursor.close()
            conn.close()


# ============================================================
# テーブル候補（<platform>.fujinp）
# ============================================================
@table_post_bp.route('/get_fujinp_tables', methods=['GET'])
@login_required
def get_fujinp_tables():
    """
    <platform>.fujinp のテーブル一覧を辞書順で返す（admin: プロジェクト割り当て用）。
    DB名はconfigの定数から取得し、platform非依存にする。
    """
    if not _is_admin():
        return jsonify({'success': False, 'error': '管理者権限が必要です'}), 403

    # テーブルサイクルと同じく、platform非依存にDB名を組み立てる
    fujinp_db = f"{Config.DB_ACCOUNT}$fujinp"
    try:
        conn = mysql.connector.connect(**DatabaseConfig.get_config(fujinp_db))
        cursor = conn.cursor()
        cursor.execute("SHOW TABLES")
        tables = sorted(r[0] for r in cursor.fetchall())
        return jsonify({'success': True, 'database': fujinp_db,
                        'tables': tables})
    except mysql.connector.Error as e:
        logging.error("get_fujinp_tables error: %s", e)
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        if 'conn' in locals() and conn.is_connected():
            cursor.close()
            conn.close()


@table_post_bp.route('/assign_tables', methods=['POST'])
@login_required
def assign_tables():
    """
    プロジェクトに対象テーブルを割り当てる（admin限定）。
    送られたテーブル集合で table_post_status を同期する。
    チェックが外れたテーブルは更新実績の有無に関わらず削除する
    （更新記録 table_post_history のスナップショットは別途残るため復元は可能）。
    """
    if not _is_admin():
        return jsonify({'success': False, 'error': '管理者権限が必要です'}), 403

    data = request.json or {}
    project_id = data.get('project_id')
    database = data.get('database')
    tables = data.get('tables') or []
    if not project_id or not database:
        return jsonify({'success': False,
                        'error': 'project_id または database が未指定です'}), 400

    try:
        conn = mysql.connector.connect(**DatabaseConfig.default())
        cursor = conn.cursor(dictionary=True)

        cursor.execute(
            f"SELECT id, table_name FROM {T_STATUS} "
            f"WHERE project_id = %s AND database_name = %s",
            (project_id, database))
        existing = {row['table_name']: row for row in cursor.fetchall()}

        wanted = set(tables)
        # 追加
        for tbl in wanted - set(existing.keys()):
            cursor.execute(
                f"INSERT INTO {T_STATUS} "
                f"(project_id, database_name, table_name) VALUES (%s, %s, %s)",
                (project_id, database, tbl))
        # 削除（チェックが外れたものは実績の有無に関わらず削除）
        for tbl, row in existing.items():
            if tbl not in wanted:
                cursor.execute(f"DELETE FROM {T_STATUS} WHERE id = %s",
                               (row['id'],))

        conn.commit()
        return jsonify({'success': True})
    except mysql.connector.Error as e:
        logging.error("assign_tables error: %s", e)
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        if 'conn' in locals() and conn.is_connected():
            cursor.close()
            conn.close()


# ============================================================
# 更新状態表示テーブル
# ============================================================
@table_post_bp.route('/get_status', methods=['GET'])
@login_required
def get_status():
    """
    プロジェクトの更新状態一覧（テーブル名の辞書順）。
    各テーブルについて、未適用(pending)の申請があれば
    その件数・最新申請者・最新申請日時も付与する。
    """
    project_id = request.args.get('project_id')
    if not project_id:
        return jsonify({'success': False,
                        'error': 'project_id が未指定です'}), 400
    try:
        conn = mysql.connector.connect(**DatabaseConfig.default())
        cursor = conn.cursor(dictionary=True)
        cursor.execute(
            f"SELECT id, project_id, database_name, table_name, assignee, "
            f"due_date, last_updater, last_updated_at, final_status "
            f"FROM {T_STATUS} WHERE project_id = %s "
            f"ORDER BY table_name ASC", (project_id,))
        rows = cursor.fetchall()

        # 未適用(pending)申請をテーブル単位で集計
        cursor.execute(
            f"SELECT database_name, table_name, COUNT(*) AS pending_count, "
            f"MAX(id) AS latest_id "
            f"FROM {T_REQUESTS} "
            f"WHERE project_id = %s AND status = 'pending' "
            f"GROUP BY database_name, table_name", (project_id,))
        pending = {(p['database_name'], p['table_name']): p
                   for p in cursor.fetchall()}

        # 最新pending申請の申請者・申請日時を引く
        latest_ids = [p['latest_id'] for p in pending.values()]
        latest_info = {}
        if latest_ids:
            placeholders = ', '.join(['%s'] * len(latest_ids))
            cursor.execute(
                f"SELECT id, submitted_by, submitted_at "
                f"FROM {T_REQUESTS} WHERE id IN ({placeholders})",
                tuple(latest_ids))
            latest_info = {r['id']: r for r in cursor.fetchall()}

        for r in rows:
            r['due_date'] = str(r['due_date']) if r['due_date'] else ''
            r['last_updated_at'] = (str(r['last_updated_at'])
                                    if r['last_updated_at'] else '')
            key = (r['database_name'], r['table_name'])
            p = pending.get(key)
            if p:
                info = latest_info.get(p['latest_id'], {})
                r['pending_count'] = p['pending_count']
                r['pending_by'] = info.get('submitted_by', '')
                r['pending_at'] = (str(info.get('submitted_at'))
                                   if info.get('submitted_at') else '')
            else:
                r['pending_count'] = 0
                r['pending_by'] = ''
                r['pending_at'] = ''

        return jsonify({'success': True, 'status': rows})
    except mysql.connector.Error as e:
        logging.error("get_status error: %s", e)
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        if 'conn' in locals() and conn.is_connected():
            cursor.close()
            conn.close()


@table_post_bp.route('/update_status_cell', methods=['POST'])
@login_required
def update_status_cell():
    """状態テーブルの assignee / due_date / final_status を書き換える（admin限定）"""
    if not _is_admin():
        return jsonify({'success': False, 'error': '管理者権限が必要です'}), 403

    data = request.json or {}
    status_id = data.get('status_id')
    field = data.get('field')
    value = data.get('value')
    if field not in ('assignee', 'due_date', 'final_status'):
        return jsonify({'success': False, 'error': '不正なフィールドです'}), 400
    if not status_id:
        return jsonify({'success': False,
                        'error': 'status_id が未指定です'}), 400

    if field == 'due_date' and not value:
        value = None
    try:
        conn = mysql.connector.connect(**DatabaseConfig.default())
        cursor = conn.cursor()
        cursor.execute(
            f"UPDATE {T_STATUS} SET {field} = %s WHERE id = %s",
            (value, status_id))
        conn.commit()
        return jsonify({'success': True})
    except mysql.connector.Error as e:
        logging.error("update_status_cell error: %s", e)
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        if 'conn' in locals() and conn.is_connected():
            cursor.close()
            conn.close()


# ============================================================
# ダウンロード（一般ユーザ可）
# ============================================================
@table_post_bp.route('/download_table', methods=['GET'])
@login_required
def download_table():
    """対象テーブルをxlsxでダウンロード（TIME型対応、テーブルサイクル準拠）"""
    database = request.args.get('database')
    table_name = request.args.get('table')
    if not database or not table_name:
        return "database or table not specified", 400

    try:
        conn = mysql.connector.connect(**DatabaseConfig.get_config(database))
        cursor = conn.cursor(dictionary=True)

        cursor.execute(f"DESCRIBE `{table_name}`")
        columns_info = cursor.fetchall()
        time_columns = {col['Field'] for col in columns_info
                        if col['Type'].lower().startswith('time')}

        cursor.execute(f"SELECT * FROM `{table_name}`")
        rows = cursor.fetchall()
        df = pd.DataFrame(rows)

        for col in time_columns:
            if col in df.columns:
                df[col] = df[col].apply(
                    lambda x: x.total_seconds() / 86400
                    if isinstance(x, (datetime.timedelta, pd.Timedelta))
                    else ((x.hour * 3600 + x.minute * 60 + x.second) / 86400
                          if isinstance(x, datetime.time) else x))

        output = io.BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name=table_name[:31], index=False)
            worksheet = writer.sheets[table_name[:31]]
            from openpyxl.utils import get_column_letter
            for idx, col_name in enumerate(df.columns):
                if col_name in time_columns:
                    col_letter = get_column_letter(idx + 1)
                    for cell in worksheet[col_letter]:
                        if cell.row > 1:
                            cell.number_format = 'h:mm:ss'
        output.seek(0)

        filename = f"{database}_{table_name}_download.xlsx"
        return send_file(
            output, as_attachment=True, download_name=filename,
            mimetype=("application/vnd.openxmlformats-officedocument"
                      ".spreadsheetml.sheet"))
    except Exception as e:
        logging.error("download_table error: %s", e)
        return str(e), 500
    finally:
        if 'conn' in locals() and conn.is_connected():
            cursor.close()
            conn.close()


# ============================================================
# 更新申請（一般ユーザ可）
# ============================================================
def _project_name(project_id):
    """プロジェクト名を引く（通知本文の表示用。失敗しても None を返すだけ）"""
    try:
        conn = mysql.connector.connect(**DatabaseConfig.default())
        cursor = conn.cursor()
        cursor.execute(f"SELECT project_name FROM {T_PROJECTS} WHERE id = %s",
                       (project_id,))
        row = cursor.fetchone()
        return row[0] if row else None
    except Exception as e:
        logging.warning("_project_name 失敗（project_id=%s）: %s",
                        project_id, e)
        return None
    finally:
        if 'conn' in locals() and conn.is_connected():
            cursor.close()
            conn.close()


def _notify_request_submitted(request_id, project_id, database,
                              table_name, row_count):
    """
    更新申請の受付を TABLE_POST_SLACK_CHANNEL へ通知する。
    notifiers（共通通知サブシステム）経由なので例外は出ず、
    送信は通知台帳 notify_ledger にも自動記録される。
    通知の失敗で申請受付（本処理）を止めない。
    """
    lines = [
        '<!channel> テーブルポストに更新申請がありました。',
        '',
        f'■ 申請番号：#{request_id}',
    ]
    pname = _project_name(project_id)
    if pname:
        lines.append(f'■ プロジェクト：{pname}')
    lines.append(f'■ 対象テーブル：{database} / {table_name}')
    lines.append(f'■ 申請者：{_current_user()}')
    lines.append(f'■ 行数：{row_count} 行')
    lines.append('')
    lines.append('管理者はテーブルポストの画面で内容を確認し、'
                 '適用または却下してください。')

    notify_channel(
        getattr(Config, 'TABLE_POST_SLACK_CHANNEL', None),
        '\n'.join(lines),
        log_label='更新申請受付通知',
        sender=_current_user(),
        app='table_post',
    )


@table_post_bp.route('/submit_request', methods=['POST'])
@login_required
def submit_request():
    """
    編集済xlsxをアップロードして更新申請を登録する。
    スキーマ（カラム名・順序）が対象テーブルと不一致なら reject。
    NOT NULL 列の空値、PRIMARY/UNIQUE 索引での重複も reject（_validate_rows）。
    内容はJSON化して table_post_requests に保存。本番テーブルは触らない。
    """
    project_id = request.form.get('project_id')
    database = request.form.get('database')
    table_name = request.form.get('table_name')
    file = request.files.get('excel_file')

    if not all([project_id, database, table_name, file]):
        return jsonify({'success': False,
                        'error': 'パラメータが不足しています'}), 400

    try:
        df = pd.read_excel(file, engine='openpyxl')
    except Exception as e:
        return jsonify({'success': False,
                        'error': f'Excel読込失敗: {e}'}), 400

    # --- スキーマ点検（カラム名と順序が完全一致すること） ---
    try:
        conn = mysql.connector.connect(**DatabaseConfig.get_config(database))
        cursor = conn.cursor()
        table_columns = _get_table_schema(cursor, table_name)
    except mysql.connector.Error as e:
        return jsonify({'success': False,
                        'error': f'テーブル参照失敗: {e}'}), 500
    finally:
        if 'conn' in locals() and conn.is_connected():
            cursor.close()
            conn.close()

    excel_columns = [str(c) for c in df.columns]
    if excel_columns != table_columns:
        return jsonify({
            'success': False,
            'error': ('スキーマ不一致のため受理できません。'
                      '列の追加・削除・並べ替えは禁止です。\n'
                      f'テーブル: {table_columns}\n'
                      f'アップロード: {excel_columns}')
        }), 400

    # --- 内容をJSON化 ---
    for col in df.columns:
        if pd.api.types.is_datetime64_any_dtype(df[col]):
            df[col] = df[col].dt.strftime('%Y-%m-%d %H:%M:%S')
    rows = [[_to_python(v) for v in row] for row in df.values]

    # --- 制約点検（NOT NULL・PRIMARY/UNIQUE 重複）。違反は受理しない ---
    try:
        rows, filled, problems = _check_payload(database, table_name,
                                                table_columns, rows)
    except mysql.connector.Error as e:
        return jsonify({'success': False,
                        'error': f'制約参照失敗: {e}'}), 500
    if problems:
        return jsonify({
            'success': False,
            'error': ('制約違反のため受理できません。xlsx を修正して'
                      '再申請してください。\n' + '\n'.join(problems))
        }), 400

    payload = json.dumps({'columns': table_columns, 'rows': rows},
                          ensure_ascii=False, default=str)

    try:
        conn = mysql.connector.connect(**DatabaseConfig.default())
        cursor = conn.cursor()
        now = get_jst_now().strftime('%Y-%m-%d %H:%M:%S')
        cursor.execute(
            f"INSERT INTO {T_REQUESTS} "
            f"(project_id, database_name, table_name, submitted_by, "
            f"submitted_at, payload, row_count, status) "
            f"VALUES (%s, %s, %s, %s, %s, %s, %s, 'pending')",
            (project_id, database, table_name, _current_user(),
             now, payload, len(rows)))
        conn.commit()
        request_id = cursor.lastrowid

        # ── Slack 通知（commit 成功後にのみ発出。成否は本処理に影響しない）──
        # 宛先は config の TABLE_POST_SLACK_CHANNEL（チャンネルID。C...）。
        # 未設定の間は notifiers が {'ok': False} を返してログと台帳に残るだけ。
        _notify_request_submitted(request_id, project_id, database,
                                  table_name, len(rows))

        note = (f'（空セル {filled} 個は NOT NULL 列のため空文字／既定値'
                f'として扱いました）' if filled else '')
        return jsonify({'success': True,
                        'message': f'{len(rows)} 行の更新申請を'
                                   f'受け付けました。{note}',
                        'request_id': request_id})
    except mysql.connector.Error as e:
        logging.error("submit_request error: %s", e)
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        if 'conn' in locals() and conn.is_connected():
            cursor.close()
            conn.close()


@table_post_bp.route('/get_requests', methods=['GET'])
@login_required
def get_requests():
    """更新申請一覧（payload本体は除く軽量版）。adminの適用画面用。"""
    if not _is_admin():
        return jsonify({'success': False, 'error': '管理者権限が必要です'}), 403

    project_id = request.args.get('project_id')
    try:
        conn = mysql.connector.connect(**DatabaseConfig.default())
        cursor = conn.cursor(dictionary=True)
        sql = (f"SELECT id, project_id, database_name, table_name, "
               f"submitted_by, submitted_at, row_count, status "
               f"FROM {T_REQUESTS}")
        params = ()
        if project_id:
            sql += " WHERE project_id = %s"
            params = (project_id,)
        sql += " ORDER BY id DESC"
        cursor.execute(sql, params)
        rows = cursor.fetchall()
        for r in rows:
            r['submitted_at'] = (str(r['submitted_at'])
                                 if r['submitted_at'] else '')
        return jsonify({'success': True, 'requests': rows})
    except mysql.connector.Error as e:
        logging.error("get_requests error: %s", e)
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        if 'conn' in locals() and conn.is_connected():
            cursor.close()
            conn.close()


# ============================================================
# 申請の適用（admin限定）
# ============================================================
@table_post_bp.route('/apply_request', methods=['POST'])
@login_required
def apply_request():
    """
    更新申請を本番テーブルに適用する（admin限定）。
    手順: 適用前内容を table_post_history に退避 → 本番をDELETE/INSERTで置換
          → 申請statusを applied → 状態テーブルを更新。
    """
    if not _is_admin():
        return jsonify({'success': False, 'error': '管理者権限が必要です'}), 403

    data = request.json or {}
    request_id = data.get('request_id')
    if not request_id:
        return jsonify({'success': False,
                        'error': 'request_id が未指定です'}), 400

    # --- 申請を取得 ---
    try:
        conn = mysql.connector.connect(**DatabaseConfig.default())
        cursor = conn.cursor(dictionary=True)
        cursor.execute(f"SELECT * FROM {T_REQUESTS} WHERE id = %s",
                       (request_id,))
        req = cursor.fetchone()
    finally:
        if 'conn' in locals() and conn.is_connected():
            cursor.close()
            conn.close()

    if not req:
        return jsonify({'success': False,
                        'error': '申請が見つかりません'}), 404
    if req['status'] == 'applied':
        return jsonify({'success': False,
                        'error': 'この申請は既に適用済みです'}), 400

    database = req['database_name']
    table_name = req['table_name']

    try:
        payload = json.loads(req['payload'])
        new_columns = payload['columns']
        new_rows = payload['rows']

        # 制約点検（点検前に登録された申請の救済。違反は適用せず理由を返す）
        new_rows, _filled, problems = _check_payload(
            database, table_name, new_columns, new_rows)
        if problems:
            return jsonify({
                'success': False,
                'error': ('制約違反のため適用できません。申請者に xlsx の'
                          '修正と再申請を依頼してください。\n'
                          + '\n'.join(problems))
            }), 400

        # 適用前の本番内容を履歴に退避
        before = _read_table_as_dict(database, table_name)
        _save_history(request_id, database, table_name, before,
                      note=f'applied request #{request_id}')

        # 本番テーブルを置換
        applied_count = _restore_table(database, table_name,
                                       new_columns, new_rows)

        # 申請statusと状態テーブルを更新
        now = get_jst_now().strftime('%Y-%m-%d %H:%M:%S')
        conn = mysql.connector.connect(**DatabaseConfig.default())
        cursor = conn.cursor()
        cursor.execute(
            f"UPDATE {T_REQUESTS} SET status = 'applied' WHERE id = %s",
            (request_id,))
        cursor.execute(
            f"UPDATE {T_STATUS} SET last_updater = %s, last_updated_at = %s, "
            f"final_status = %s "
            f"WHERE project_id = %s AND database_name = %s AND table_name = %s",
            (req['submitted_by'], now, '正常終了',
             req['project_id'], database, table_name))
        conn.commit()

        # ── 公式データ集への由来記録（対象が公式テーブルの場合のみ） ──
        archived = _archive_record(
            database, table_name,
            source_ref=f'table_post:{request_id}',
            note=(f'テーブルポストで承認・適用'
                  f'（申請 #{request_id}、{applied_count}行）'),
            uploaded_by_name=req['submitted_by'],
            uploaded_at=req['submitted_at'])
        suffix = '公式データ集にも記録しました。' if archived else ''
        return jsonify({'success': True,
                        'message': f'{applied_count} 行を本番テーブルに'
                                   f'適用しました。' + suffix})
    except Exception as e:
        logging.error("apply_request error: %s", e)
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        if 'conn' in locals() and conn.is_connected():
            cursor.close()
            conn.close()


@table_post_bp.route('/reject_request', methods=['POST'])
@login_required
def reject_request():
    """申請を却下する（admin限定、本番テーブルは触らない）"""
    if not _is_admin():
        return jsonify({'success': False, 'error': '管理者権限が必要です'}), 403

    data = request.json or {}
    request_id = data.get('request_id')
    if not request_id:
        return jsonify({'success': False,
                        'error': 'request_id が未指定です'}), 400
    try:
        conn = mysql.connector.connect(**DatabaseConfig.default())
        cursor = conn.cursor()
        cursor.execute(
            f"UPDATE {T_REQUESTS} SET status = 'rejected' "
            f"WHERE id = %s AND status = 'pending'", (request_id,))
        conn.commit()
        return jsonify({'success': True})
    except mysql.connector.Error as e:
        logging.error("reject_request error: %s", e)
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        if 'conn' in locals() and conn.is_connected():
            cursor.close()
            conn.close()


# ============================================================
# 更新記録 / ロールバック（admin限定）
# ============================================================
def _save_history(request_id, database, table_name, snapshot, note):
    """更新記録テーブルに版を保存"""
    conn = mysql.connector.connect(**DatabaseConfig.default())
    try:
        cursor = conn.cursor()
        now = get_jst_now().strftime('%Y-%m-%d %H:%M:%S')
        cursor.execute(
            f"INSERT INTO {T_HISTORY} "
            f"(request_id, database_name, table_name, recorded_at, "
            f"snapshot, row_count, note) VALUES (%s, %s, %s, %s, %s, %s, %s)",
            (request_id, database, table_name, now,
             json.dumps(snapshot, ensure_ascii=False, default=str),
             len(snapshot.get('rows', [])), note))
        conn.commit()
        return cursor.lastrowid
    finally:
        if conn.is_connected():
            cursor.close()
            conn.close()


@table_post_bp.route('/get_history', methods=['GET'])
@login_required
def get_history():
    """更新記録一覧（snapshot本体は除く軽量版）。admin限定。"""
    if not _is_admin():
        return jsonify({'success': False, 'error': '管理者権限が必要です'}), 403

    database = request.args.get('database')
    table_name = request.args.get('table')
    try:
        conn = mysql.connector.connect(**DatabaseConfig.default())
        cursor = conn.cursor(dictionary=True)
        sql = (f"SELECT id, request_id, database_name, table_name, "
               f"recorded_at, row_count, note FROM {T_HISTORY}")
        params = ()
        if database and table_name:
            sql += " WHERE database_name = %s AND table_name = %s"
            params = (database, table_name)
        sql += " ORDER BY id DESC"
        cursor.execute(sql, params)
        rows = cursor.fetchall()
        for r in rows:
            r['recorded_at'] = (str(r['recorded_at'])
                                if r['recorded_at'] else '')
        return jsonify({'success': True, 'history': rows})
    except mysql.connector.Error as e:
        logging.error("get_history error: %s", e)
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        if 'conn' in locals() and conn.is_connected():
            cursor.close()
            conn.close()


@table_post_bp.route('/rollback', methods=['POST'])
@login_required
def rollback():
    """
    更新記録テーブルの指定の版を本番テーブルに復元する（admin限定）。
    復元前の現状もまた履歴に退避するので、復元操作自体も巻き戻せる。
    """
    if not _is_admin():
        return jsonify({'success': False, 'error': '管理者権限が必要です'}), 403

    data = request.json or {}
    history_id = data.get('history_id')
    if not history_id:
        return jsonify({'success': False,
                        'error': 'history_id が未指定です'}), 400

    # 復元元の版を取得
    try:
        conn = mysql.connector.connect(**DatabaseConfig.default())
        cursor = conn.cursor(dictionary=True)
        cursor.execute(f"SELECT * FROM {T_HISTORY} WHERE id = %s",
                       (history_id,))
        hist = cursor.fetchone()
    finally:
        if 'conn' in locals() and conn.is_connected():
            cursor.close()
            conn.close()

    if not hist:
        return jsonify({'success': False,
                        'error': '記録が見つかりません'}), 404

    database = hist['database_name']
    table_name = hist['table_name']
    try:
        snapshot = json.loads(hist['snapshot'])

        # 復元前の現状を退避（復元操作も巻き戻せるように）
        before = _read_table_as_dict(database, table_name)
        before_id = _save_history(
            None, database, table_name, before,
            note=f'before rollback to history #{history_id}')

        # 指定の版に復元
        count = _restore_table(database, table_name,
                               snapshot['columns'], snapshot['rows'])

        # ── 公式データ集への由来記録（対象が公式テーブルの場合のみ）。
        #    退避レコードid（ロールバック1回ごとに一意）を照合キーにする。 ──
        archived = _archive_record(
            database, table_name,
            source_ref=f'table_post:rb:{before_id}',
            note=(f'テーブルポストでロールバック'
                  f'（記録 #{history_id} の版に復元、{count}行）'),
            uploaded_by_name=_current_user())
        suffix = '公式データ集にも記録しました。' if archived else ''
        return jsonify({'success': True,
                        'message': f'記録 #{history_id} の内容（{count} 行）'
                                   f'を本番テーブルに復元しました。' + suffix})
    except Exception as e:
        logging.error("rollback error: %s", e)
        return jsonify({'success': False, 'error': str(e)}), 500


# ============================================================
# 公式データ集への記録マイグレート（admin限定）
#   承認済み更新（status='applied'、過去分を含む）の「由来の記録」を
#   official_data_archive_updates へ取り込む。テーブル本体は触れない。
#   重複点検: source_ref = 'table_post:<申請id>' で既存照合し、
#   取り込み済みはスキップ（何度実行しても安全）。
# ============================================================
ARCHIVE_TABLES_T  = 'official_data_archive_tables'
ARCHIVE_UPDATES_T = 'official_data_archive_updates'


def _logical_db_name(name):
    """'<アカウント>$fujinp' → 'fujinp' のように論理DB名へ変換する。"""
    if name and '$' in name:
        return name.rsplit('$', 1)[1]
    return name or ''


def _resolve_user_id(cursor, name):
    """
    table_post の記録（アカウント名/氏名/数値ID文字列）から users.id を引く。
    解決できなければ None（名前は別カラムに保存されるので情報は失わない）。
    """
    if name is None:
        return None
    s = str(name).strip()
    if not s:
        return None
    if s.isdigit():
        cursor.execute("SELECT id FROM users WHERE id = %s", (int(s),))
        r = cursor.fetchone()
        if r:
            return r['id']
    for col in ('username', 'full_name', 'email'):
        try:
            cursor.execute(
                f"SELECT id FROM users WHERE {col} = %s LIMIT 1", (s,))
            r = cursor.fetchone()
            if r:
                return r['id']
        except mysql.connector.Error:
            continue   # その列が users に存在しない場合は次の候補へ
    return None


def _archive_record(database_name, table_name, source_ref, note,
                    uploaded_by_name=None, uploaded_at=None):
    """
    公式テーブル（official_data_archive の登録簿に載っているテーブル）を
    変更したとき、公式データ集の更新履歴へ由来の記録を1行追加する。
      - 対象が登録簿に無ければ何もしない（公式テーブル以外は対象外）
      - source_ref のUNIQUE制約により二重記録は自動的に防がれる
      - 記録の失敗で本処理（適用・ロールバック）を止めない。例外は投げない。
    戻り値: 記録できたら True、対象外・重複・失敗は False。
    """
    try:
        ldb = _logical_db_name(database_name)
        now = get_jst_now()
        conn = mysql.connector.connect(**DatabaseConfig.default())
        cursor = conn.cursor(dictionary=True)
        cursor.execute(
            f"""SELECT id FROM {ARCHIVE_TABLES_T}
                WHERE database_name = %s AND table_name = %s""",
            (ldb, table_name))
        item = cursor.fetchone()
        if not item:
            return False   # 公式テーブルではない（記録対象外）
        uploaded_by = _resolve_user_id(cursor, uploaded_by_name)
        try:
            applied_by = int(session.get('user_id'))
        except (TypeError, ValueError):
            applied_by = None
        cursor.execute(
            f"""INSERT INTO {ARCHIVE_UPDATES_T}
                (table_item_id, table_name, database_name,
                 original_filename, stored_filename, note,
                 uploaded_by, uploaded_by_name, uploaded_at,
                 status, applied_by, applied_at, source_ref)
                VALUES (%s, %s, %s, NULL, %s, %s,
                        %s, %s, %s, 'applied', %s, %s, %s)""",
            (item['id'], table_name, ldb,
             'migrated:' + source_ref, note,
             uploaded_by,
             (str(uploaded_by_name) if uploaded_by_name else None),
             uploaded_at or now,
             applied_by, now, source_ref))
        conn.commit()
        return True
    except mysql.connector.IntegrityError:
        return False   # source_ref 重複＝記録済み。問題なし。
    except Exception as e:
        logging.warning("_archive_record 失敗（%s, %s）: %s",
                        table_name, source_ref, e)
        return False
    finally:
        if 'conn' in locals() and conn.is_connected():
            cursor.close()
            conn.close()


@table_post_bp.route('/migrate_to_archive', methods=['POST'])
@login_required
def migrate_to_archive():
    """
    選択中プロジェクトの承認済み更新の記録を公式データ集へ取り込む。
    リクエスト: { "project_id": ... }
    マッピング:
        uploaded_by / uploaded_by_name / uploaded_at ← 申請者・申請日時
        applied_at ← table_post_history の request_id 付き退避の記録日時
                     （なければ申請日時で代用）
        applied_by ← NULL（table_post は適用者の身元を記録していないため）
        status     ← 'applied'
        source_ref ← 'table_post:<申請id>'（重複点検キー）
    """
    if not _is_admin():
        return jsonify({'success': False, 'error': '管理者権限が必要です'}), 403
    data = request.json or {}
    project_id = data.get('project_id')
    if not project_id:
        return jsonify({'success': False,
                        'error': 'project_id が未指定です'}), 400
    try:
        conn = mysql.connector.connect(**DatabaseConfig.default())
        cursor = conn.cursor(dictionary=True)

        # ── 承認済み申請（過去分を含む） ──
        cursor.execute(
            f"""SELECT id, database_name, table_name, submitted_by,
                       submitted_at, row_count
                FROM {T_REQUESTS}
                WHERE project_id = %s AND status = 'applied'
                ORDER BY id""", (project_id,))
        reqs = cursor.fetchall()

        # ── 取り込み済みの照合キー ──
        cursor.execute(
            f"""SELECT source_ref FROM {ARCHIVE_UPDATES_T}
                WHERE source_ref LIKE 'table_post:%'""")
        existing = {r['source_ref'] for r in cursor.fetchall()}

        # ── 公式テーブル登録簿（(論理DB, テーブル名) → アイテムid） ──
        cursor.execute(
            f"SELECT id, table_name, database_name FROM {ARCHIVE_TABLES_T}")
        items = {(r['database_name'], r['table_name']): r['id']
                 for r in cursor.fetchall()}

        # ── 適用日時（request_id 付き退避レコードの記録日時） ──
        cursor.execute(
            f"""SELECT request_id, MIN(recorded_at) AS applied_at
                FROM {T_HISTORY}
                WHERE request_id IS NOT NULL
                GROUP BY request_id""")
        applied_at_map = {r['request_id']: r['applied_at']
                          for r in cursor.fetchall()}

        migrated = 0
        skipped_dup = 0
        unregistered = []
        unresolved = []
        user_cache = {}

        for rq in reqs:
            ref = f"table_post:{rq['id']}"
            if ref in existing:
                skipped_dup += 1
                continue
            ldb = _logical_db_name(rq['database_name'])
            item_id = items.get((ldb, rq['table_name']))
            if not item_id:
                label = f"{ldb} / {rq['table_name']}"
                if label not in unregistered:
                    unregistered.append(label)
                continue
            name = rq['submitted_by']
            if name not in user_cache:
                user_cache[name] = _resolve_user_id(cursor, name)
                if user_cache[name] is None and name:
                    unresolved.append(str(name))
            uid = user_cache[name]
            applied_at = applied_at_map.get(rq['id']) or rq['submitted_at']
            note = ('テーブルポストから移行（申請 #%s%s）'
                    % (rq['id'],
                       '、%s行' % rq['row_count']
                       if rq.get('row_count') is not None else ''))
            cursor.execute(
                f"""INSERT INTO {ARCHIVE_UPDATES_T}
                    (table_item_id, table_name, database_name,
                     original_filename, stored_filename, note,
                     uploaded_by, uploaded_by_name, uploaded_at,
                     status, applied_by, applied_at, source_ref)
                    VALUES (%s, %s, %s, NULL, %s, %s,
                            %s, %s, %s, 'applied', NULL, %s, %s)""",
                (item_id, rq['table_name'], ldb,
                 'migrated:' + ref, note,
                 uid, (str(name) if name else None), rq['submitted_at'],
                 applied_at, ref))
            migrated += 1

        conn.commit()
        return jsonify({
            'success': True,
            'total_applied_requests': len(reqs),
            'migrated': migrated,
            'skipped_duplicate': skipped_dup,
            'skipped_unregistered': len(reqs) - migrated - skipped_dup,
            'unregistered_tables': unregistered,
            'unresolved_users': sorted(set(unresolved)),
        })
    except Exception as e:
        logging.error("migrate_to_archive error: %s", e)
        if 'conn' in locals():
            conn.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        if 'conn' in locals() and conn.is_connected():
            cursor.close()
            conn.close()


@table_post_bp.route('/return_to_fujin')
@login_required
def return_to_fujin():
    """FUJIN-Pダッシュボードに戻る"""
    return redirect_to_dashboard()