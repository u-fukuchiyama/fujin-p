# SPDX-FileCopyrightText: 2024-2026 Toyoaki Nishida
# SPDX-License-Identifier: AGPL-3.0-or-later
#
# This file is part of FUJIN-P.
# Copyright (C) 2024-2026 Toyoaki Nishida
#
# FUJIN-P is free software: you can redistribute it and/or modify
# it under the terms of the GNU Affero General Public License as published by
# the Free Software Foundation, either version 3 of the License, or
# (at your option) any later version.
#
# FUJIN-P is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
# GNU Affero General Public License for more details.
#
# You should have received a copy of the GNU Affero General Public License
# along with FUJIN-P.  If not, see <https://www.gnu.org/licenses/>.
#
# Source: https://github.com/u-fukuchiyama/fujin-p

# table_cycle.py
# 【新プラットフォーム対応版 - 最終修正版】

import io
import json
import base64
import datetime
import logging
import mysql.connector
import pandas as pd
from decimal import Decimal
from flask import (Blueprint, request, jsonify, send_file, session,
                   render_template, Response, flash, redirect, url_for)
from decorators import login_required
from config import Config
from auth import redirect_to_dashboard
from db import DatabaseConfig, Tables
from . import table_cycle_bp
from urllib.parse import quote
import gzip


logging.basicConfig(level=logging.DEBUG)

# 監査テーブル名
# SNAPSHOT_TABLE = "table_snapshots"

SNAPSHOT_TABLE = Tables.TABLE_SNAPSHOTS

# 日時はJSTで扱う（FUJIN-Pの日時3層ルール：DBにはJSTのDATETIME，
# バックエンドで文字列化，フロントは文字列のみ）．
# サーバのローカル時刻がUTCでも正しくJSTになるよう，タイムゾーン付きで取得する．
JST = datetime.timezone(datetime.timedelta(hours=9), 'JST')


def now_jst_str(fmt='%Y-%m-%d %H:%M:%S'):
    """現在時刻をJSTの文字列で返す"""
    return datetime.datetime.now(JST).strftime(fmt)


# ============================================
# アクセス制御（アプリ全体を管理者に限定）
# ============================================

@table_cycle_bp.before_request
def require_admin():
    """
    table_cycle 配下の全エンドポイントを管理者に限定する．

    判定内容は decorators.admin_required と同一だが，本アプリの画面は
    jQuery の Ajax でサーバと通信するため，拒否時に HTML へリダイレクトすると
    画面側が応答を解釈できない．そこで Ajax には JSON を返し，
    通常のページ遷移だけリダイレクトするよう before_request で処理する．
    """
    wants_json = (request.method == 'POST'
                  or request.headers.get('X-Requested-With') == 'XMLHttpRequest')

    if 'user_id' not in session:
        if wants_json:
            return jsonify({'success': False, 'error': 'ログインが必要です'}), 401
        flash('ログインが必要です', 'error')
        return redirect(url_for('auth.login', next=request.url))

    if session.get('user_category') != 'admin':
        if wants_json:
            return jsonify({'success': False, 'error': '管理者権限が必要です'}), 403
        flash('管理者権限が必要です', 'error')
        return redirect(url_for('guest.dashboard'))

    return None


BATCH_SIZE = 20

def _encode_value(v):
    """MySQLの値をJSON安全な形にエンコード（型を失わずに）"""
    if v is None:
        return None
    if isinstance(v, (bytes, bytearray)):
        return {'__t': 'b64', 'v': base64.b64encode(v).decode('ascii')}
    if isinstance(v, Decimal):
        return {'__t': 'dec', 'v': str(v)}
    if isinstance(v, datetime.datetime):
        return {'__t': 'dt', 'v': v.strftime('%Y-%m-%d %H:%M:%S.%f')}
    if isinstance(v, datetime.date):
        return {'__t': 'date', 'v': v.isoformat()}
    if isinstance(v, datetime.timedelta):     # MySQL TIME型は timedelta で返る
        return {'__t': 'time', 'v': v.total_seconds()}
    if isinstance(v, datetime.time):
        return {'__t': 'time2', 'v': v.isoformat()}
    return v                                  # int, float, bool, str はそのまま


def _decode_value(v):
    """エンコードされた値を元のPython型へ復元"""
    if isinstance(v, dict) and '__t' in v:
        t = v['__t']
        if t == 'b64':
            return base64.b64decode(v['v'])
        if t == 'dec':
            return Decimal(v['v'])
        if t == 'dt':
            return datetime.datetime.strptime(v['v'], '%Y-%m-%d %H:%M:%S.%f')
        if t == 'date':
            return datetime.date.fromisoformat(v['v'])
        if t == 'time':
            return datetime.timedelta(seconds=v['v'])
        if t == 'time2':
            return datetime.time.fromisoformat(v['v'])
    return v

@table_cycle_bp.route('/')
@login_required
def table_cycle_dashboard():
    """
    Table Cycleメイン画面
    """
    return render_template('table_cycle_dashboard.html')

@table_cycle_bp.route('/get_databases', methods=['GET'])
@login_required
def get_databases():
    """
    DB一覧を返すエンドポイント
    """
    try:
        conn = mysql.connector.connect(**DatabaseConfig.base())
        cursor = conn.cursor()  # ★ここはdictionary不要（SCHEMA_NAMEだけ取る）

        # PythonAnywhereのユーザー名プレフィックスでフィルタリング
        user_prefix = Config.DB_ACCOUNT + "$"  # ★修正：ハードコード削除
        query = f"SELECT SCHEMA_NAME FROM INFORMATION_SCHEMA.SCHEMATA WHERE SCHEMA_NAME LIKE '{user_prefix}%'"
        cursor.execute(query)
        rows = cursor.fetchall()
        dbs = [r[0] for r in rows]

        return jsonify({'success': True, 'databases': dbs})
    except mysql.connector.Error as e:
        logging.error("get_databases error: %s", e)
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        if 'conn' in locals() and conn.is_connected():
            cursor.close()
            conn.close()

@table_cycle_bp.route('/get_tables', methods=['POST'])
@login_required
def get_tables():
    """
    テーブル一覧を返すエンドポイント
    """
    data = request.json
    database = data.get('database')
    if not database:
        return jsonify({'success': False, 'error': 'database not specified'}), 400

    try:
        conn = mysql.connector.connect(**DatabaseConfig.get_config(database))
        cursor = conn.cursor()  # ★ここもdictionary不要
        cursor.execute("SHOW TABLES")
        rows = cursor.fetchall()
        tables = [r[0] for r in rows]

        return jsonify({'success': True, 'tables': tables})
    except mysql.connector.Error as e:
        logging.error("get_tables error: %s", e)
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        if 'conn' in locals() and conn.is_connected():
            cursor.close()
            conn.close()



@table_cycle_bp.route('/download_table', methods=['GET'])
@login_required
def download_table():
    """
    テーブルデータをExcelでダウンロード（TIME型完全対応版）
    """
    database = request.args.get('database')
    table_name = request.args.get('table')

    if not database or not table_name:
        return "database or table not specified", 400

    try:
        conn = mysql.connector.connect(**DatabaseConfig.get_config(database))
        cursor = conn.cursor(dictionary=True)  # ★dictionary=True に修正

        # TIME型カラムを識別
        cursor.execute(f"DESCRIBE `{table_name}`")
        columns_info = cursor.fetchall()
        time_columns = {col['Field'] for col in columns_info
                       if col['Type'].lower().startswith('time')}

        # データ取得
        cursor.execute(f"SELECT * FROM `{table_name}`")
        rows = cursor.fetchall()
        df = pd.DataFrame(rows)

        # TIME型をExcelシリアル値に変換
        for col in time_columns:
            if col in df.columns:
                df[col] = df[col].apply(lambda x:
                    x.total_seconds() / 86400 if isinstance(x, (datetime.timedelta, pd.Timedelta))
                    else ((x.hour * 3600 + x.minute * 60 + x.second) / 86400 if isinstance(x, datetime.time)
                    else x)
                )

        # Excel生成
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name=table_name[:31], index=False)

            workbook = writer.book
            worksheet = writer.sheets[table_name[:31]]

            from openpyxl.utils import get_column_letter

            # TIME型カラムに書式設定
            for idx, col_name in enumerate(df.columns):
                if col_name in time_columns:
                    col_letter = get_column_letter(idx + 1)
                    for cell in worksheet[col_letter]:
                        if cell.row > 1:
                            cell.number_format = 'h:mm:ss'

        output.seek(0)

        # スナップショット保存
        snapshot_data = {
            'database': database,
            'table_name': table_name,
            'rows': rows,
            'user_id': session.get('user_id'),
            'row_count': len(rows)
        }
        save_download_snapshot(database, table_name, snapshot_data)

        filename = f"{database}_{table_name}_download.xlsx"
        return send_file(
            output,
            as_attachment=True,
            download_name=filename,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

    except Exception as e:
        logging.error("🔍download_table Error: %s", e)
        return str(e), 500
    finally:
        if 'conn' in locals() and conn.is_connected():
            cursor.close()
            conn.close()

@table_cycle_bp.route('/download_jsonl', methods=['GET'])
@login_required
def download_jsonl():
    """
    テーブルを JSON Lines (.jsonl) でフルバックアップ
    （構造込み・ストリーミング・完全復元対応）
    """
    database = request.args.get('database')
    table_name = request.args.get('table')

    if not database or not table_name:
        return "database or table not specified", 400

    def generate():
        conn = None
        cursor = None
        try:
            conn = mysql.connector.connect(**DatabaseConfig.get_config(database))

            meta_cursor = conn.cursor(dictionary=True)
            meta_cursor.execute(f"DESCRIBE `{table_name}`")
            columns_info = meta_cursor.fetchall()
            columns = [c['Field'] for c in columns_info]

            # SHOW CREATE TABLE で DDL を取得（テーブルごと再生成用）
            meta_cursor.execute(f"SHOW CREATE TABLE `{table_name}`")
            ddl_row = meta_cursor.fetchone()
            create_sql = ddl_row.get('Create Table') or ddl_row.get('Create View')
            meta_cursor.close()

            meta = {
                '_meta': True,
                'version': 2,
                'database': database,
                'table_name': table_name,
                'columns': columns,
                'schema': [
                    {'name': c['Field'], 'type': c['Type'],
                     'null': c['Null'], 'key': c['Key'], 'default': c['Default']}
                    for c in columns_info
                ],
                'create_sql': create_sql,
                'exported_at': now_jst_str(),   # JST
            }
            yield json.dumps(meta, default=str, ensure_ascii=False) + "\n"

            # データ本体を1行ずつ（サーバーサイドカーソルでメモリに載せない）
            cursor = conn.cursor(dictionary=True, buffered=False)
            cursor.execute(f"SELECT * FROM `{table_name}`")
            for row in cursor:
                encoded = {k: _encode_value(v) for k, v in row.items()}
                yield json.dumps(encoded, ensure_ascii=False) + "\n"

        except Exception as e:
            logging.error("download_jsonl error: %s", e)
            yield json.dumps({'_error': str(e)}, ensure_ascii=False) + "\n"
        finally:
            if cursor:
                cursor.close()
            if conn and conn.is_connected():
                conn.close()

    def generate_gz():
        # gzip ストリームを少しずつ生成（メモリに全体を載せない）
        buf = io.BytesIO()
        gz = gzip.GzipFile(fileobj=buf, mode='wb')
        for chunk in generate():                 # 既存の generate() がテキスト行を yield
            gz.write(chunk.encode('utf-8'))
            gz.flush()
            data = buf.getvalue()
            if data:
                yield data
                buf.seek(0)
                buf.truncate(0)
        gz.close()
        data = buf.getvalue()                    # 残りを吐き出す
        if data:
            yield data

    filename = f"{database}_{table_name}.jsonl.gz"
    ascii_fallback = "table_export.jsonl.gz"
    encoded = quote(filename)
    disposition = f"attachment; filename=\"{ascii_fallback}\"; filename*=UTF-8''{encoded}"
    return Response(
        generate_gz(),
        mimetype='application/gzip',
        headers={'Content-Disposition': disposition}
    )

@table_cycle_bp.route('/upload_table', methods=['POST'])
@login_required
def upload_table():
    database = request.form.get('database')
    table_name = request.form.get('table_name')
    file = request.files.get('excel_file')

    if not all([database, table_name, file]):
        return jsonify({'success': False, 'error': 'Parameters missing'}), 400

    try:
        df = pd.read_excel(file, engine='openpyxl')
        logging.debug(f"Read Excel file: {df.shape} rows, {list(df.columns)} columns")

        # datetime列を変換
        for col in df.columns:
            if pd.api.types.is_datetime64_any_dtype(df[col]):
                df[col] = df[col].dt.strftime('%Y-%m-%d %H:%M:%S').replace('NaT', None)

        # ★追加：numpy/pandas型をPython標準型に変換
        def convert_value(v):
            if pd.isna(v) if not isinstance(v, str) else False:
                return None
            if hasattr(v, 'item'):          # numpy.int64, numpy.float64 など
                return v.item()
            return v

        # 以下は既存コードのまま、valuesの生成部分だけ変更
        conn = mysql.connector.connect(**DatabaseConfig.get_config(database))
        cursor = conn.cursor()

        try:
            cursor.execute("SET FOREIGN_KEY_CHECKS = 0")
            conn.start_transaction()
            cursor.execute(f"DELETE FROM `{table_name}`")

            if not df.empty:
                placeholders = ', '.join(['%s'] * len(df.columns))
                columns = '`, `'.join(str(c) for c in df.columns)
                insert_query = f"INSERT INTO `{table_name}` (`{columns}`) VALUES ({placeholders})"

                # ★変更：convert_value を使用
                values = [tuple(convert_value(v) for v in row) for row in df.values]
                cursor.executemany(insert_query, values)

            cursor.execute("SET FOREIGN_KEY_CHECKS = 1")

            snapshot_data = {
                'database': database,
                'table_name': table_name,
                'rows': df.to_dict('records'),
                'user_id': session.get('user_id'),
                'row_count': len(df)
            }
            save_upload_snapshot(database, table_name, snapshot_data)

            conn.commit()
            return jsonify({'success': True, 'message': f"{len(df)} 行インポートしました。"})

        except Exception as e:
            conn.rollback()
            cursor.execute("SET FOREIGN_KEY_CHECKS = 1")
            logging.error(f"Error during database operation: {str(e)}")
            return jsonify({'success': False, 'error': str(e)}), 500

    except Exception as e:
        logging.error(f"Error processing Excel file: {str(e)}")
        return jsonify({'success': False, 'error': str(e)}), 500

    finally:
        if 'conn' in locals() and conn.is_connected():
            cursor.close()
            conn.close()

@table_cycle_bp.route('/upload_jsonl', methods=['POST'])
@login_required
def upload_jsonl():
    """
    JSON Lines (.jsonl) フルバックアップから復元（更地前提・スキーマ自動作成）。
    - 行き先テーブル名はフォーム table_name 優先、未指定ならファイルの元テーブル名を流用。
    - 同名テーブルがあれば DROP し、保存時の構造（create_sql）で作り直してから INSERT。
    - 行き先DBは「ログイン中プラットフォーム」＋「db-select で選んだDB」で決まる。
    """
    database = request.form.get('database')
    table_name = request.form.get('table_name')   # 空ならメタ行から補完
    file = request.files.get('jsonl_file')

    if not database or not file:
        return jsonify({'success': False, 'error': 'database または file がありません'}), 400

    conn = None
    cursor = None
    try:
        conn = mysql.connector.connect(**DatabaseConfig.get_config(database))
        cursor = conn.cursor()
        cursor.execute("SET FOREIGN_KEY_CHECKS = 0")

        columns = None
        insert_q = None
        schema_ready = False
        batch = []
        total = 0

        # 先頭2バイトで gzip 判定（0x1f 0x8b なら gzip）
        head = file.stream.read(2)
        file.stream.seek(0)
        if head == b'\x1f\x8b':
            raw = gzip.GzipFile(fileobj=file.stream, mode='rb')
            stream = io.TextIOWrapper(raw, encoding='utf-8')
        else:
            stream = io.TextIOWrapper(file.stream, encoding='utf-8')

        for line in stream:
            line = line.strip()
            if not line:
                continue
            rec = json.loads(line)

            # メタ行：テーブル名を確定し、構造を準備（DROP → CREATE）
            if rec.get('_meta'):
                columns = rec.get('columns')
                create_sql = rec.get('create_sql')

                if not table_name:                      # フォーム未指定なら
                    table_name = rec.get('table_name')  # ファイルの元テーブル名を採用
                if not table_name:
                    raise ValueError('行き先テーブル名が決まりません（フォーム・ファイルとも空）')
                if not create_sql:
                    raise ValueError('create_sql がファイルに無いため復元できません')

                cursor.execute(f"DROP TABLE IF EXISTS `{table_name}`")
                cursor.execute(create_sql)
                schema_ready = True
                continue

            if rec.get('_error'):
                raise ValueError(f"エクスポートファイルにエラー行が含まれています: {rec['_error']}")

            if not schema_ready:
                raise ValueError('メタ行（_meta）が先頭にありません。フルバックアップ形式のファイルが必要です。')

            if insert_q is None:
                col_sql = '`, `'.join(columns)
                ph = ', '.join(['%s'] * len(columns))
                insert_q = f"INSERT INTO `{table_name}` (`{col_sql}`) VALUES ({ph})"

            batch.append(tuple(_decode_value(rec.get(c)) for c in columns))
            if len(batch) >= BATCH_SIZE:
                cursor.executemany(insert_q, batch)
                total += len(batch)
                batch = []

        if batch:
            cursor.executemany(insert_q, batch)
            total += len(batch)

        cursor.execute("SET FOREIGN_KEY_CHECKS = 1")
        conn.commit()
        return jsonify({'success': True, 'message': f"`{table_name}` に {total} 行を復元しました。"})

    except Exception as e:
        if conn:
            try:
                conn.rollback()
            except Exception as rollback_err:
                logging.error("rollback failed (接続切断の可能性): %s", rollback_err)
            try:
                cursor.execute("SET FOREIGN_KEY_CHECKS = 1")
            except Exception:
                pass
        logging.error("upload_jsonl error: %s", e)
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        if cursor:
            cursor.close()
        if conn and conn.is_connected():
            conn.close()


def save_download_snapshot(database, table_name, snapshot_data):
    """
    監査テーブルにダウンロード記録を保存
    """
    try:
        # conn = mysql.connector.connect(**default_db_config)
        conn = mysql.connector.connect(**DatabaseConfig.default())
        cursor = conn.cursor()
        now_str = now_jst_str()
        snapshot_json = json.dumps(snapshot_data, default=str)
        query = f"""
            INSERT INTO {SNAPSHOT_TABLE}
            (database_name, table_name, download_timestamp, download_snapshot)
            VALUES (%s, %s, %s, %s)
        """
        cursor.execute(query, (database, table_name, now_str, snapshot_json))
        conn.commit()
    except Exception as e:
        logging.error("save_download_snapshot error: %s", e)
        if conn:
            conn.rollback()
    finally:
        if 'conn' in locals() and conn.is_connected():
            cursor.close()
            conn.close()

def save_upload_snapshot(database, table_name, snapshot_data):
    """
    監査テーブルにアップロード記録を保存
    """
    try:
        # conn = mysql.connector.connect(**default_db_config)
        conn = mysql.connector.connect(**DatabaseConfig.default())
        cursor = conn.cursor()
        now_str = now_jst_str()
        snapshot_json = json.dumps(snapshot_data, default=str)
        query = f"""
            INSERT INTO {SNAPSHOT_TABLE}
            (database_name, table_name, upload_timestamp, upload_snapshot)
            VALUES (%s, %s, %s, %s)
        """
        cursor.execute(query, (database, table_name, now_str, snapshot_json))
        conn.commit()
    except Exception as e:
        logging.error("save_upload_snapshot error: %s", e)
        if conn:
            conn.rollback()
    finally:
        if 'conn' in locals() and conn.is_connected():
            cursor.close()
            conn.close()


@table_cycle_bp.route('/return_to_fujin')
@login_required
def return_to_fujin():
    """FUJIN-Pダッシュボードに戻る"""
    return redirect_to_dashboard()