# SPDX-FileCopyrightText: 2024-2026 Toyoaki Nishida
# SPDX-License-Identifier: AGPL-3.0-or-later
#
# This file is part of FUJIN-P.
# Copyright (C) 2024-2026 Toyoaki Nishida
#
# FUJIN-P is free software: you can redistribute it and/or modify
# it under the terms of the GNU Affero General Public License as published by
# the Free Software Foundation, either version 3 of the License, or
# (at your option) any later version.
#
# FUJIN-P is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
# GNU Affero General Public License for more details.
#
# You should have received a copy of the GNU Affero General Public License
# along with FUJIN-P.  If not, see <https://www.gnu.org/licenses/>.
#
# Source: https://github.com/u-fukuchiyama/fujin-p

"""
user_groups/pids.py — 永続ID（サイトをまたぐアカウント・グループの識別子）

FUJIN-P のサイトをまたいでアカウントとグループを突合するための層．
各サイトのローカルID（users.id / user_groups.id）はそのまま残し，その上に
永続IDを被せる．サイトAのローカルID 1 とサイトBのローカルID 3 が同じ永続IDを
持てば同一のアカウント，という形にする．

  ・永続IDが同定するのはアカウントであって人ではない（同じ人の複数アカウントは別ID）
  ・同一サイト内で複数行が同じ永続IDを指すことはない
  ・発番するのは源泉サイトだけ（config.py の PID_ISSUER で切り替える）
  ・配布は XLSX．受け取った側は写し取るだけで，常時同期はしない
  ・取り込みは追加のみ．配布物に載っていない所属を「消えた」とはみなさない
  ・ユーザ区分は受け入れ側を優先する．新規に作る人だけ，admin は regular に格下げする

テーブル（default DB）
  ug_pids         永続IDとローカルIDの対応（源泉サイトでは発行台帳を兼ねる）
  ug_pid_imports  取り込んだ配布物の版の記録

config.py に置く定数（どちらも無ければ既定値で動く）
  PID_ISSUER = True/False   源泉サイトなら True（発番ボタンが出る）
  PID_PREFIX = 'FUC'        発番元の識別子．永続IDの接頭辞になる
"""
import io
import re
import logging
from datetime import datetime, timedelta, timezone

from flask import render_template_string, jsonify, request, session, send_file

from decorators import login_required
from . import user_groups_bp
from .routes import get_db, check_is_total_admin

JST = timezone(timedelta(hours=9), 'JST')

KIND_TAG = {'person': 'P', 'group': 'G'}
PID_RE = re.compile(r'^[A-Z0-9]{2,8}-[PG]-\d{4,9}$')

SHEET_INFO = '情報'
SHEET_PERSON = '人'
SHEET_GROUP = 'グループ'
SHEET_MEMBER = '所属'

INFO_HEADERS = ['項目', '値']
PERSON_HEADERS = ['永続ID', 'メール', '氏名', '区分', '所属', '備考']
GROUP_HEADERS = ['永続ID', 'グループ名', '説明', '備考']
MEMBER_HEADERS = ['人の永続ID', 'グループの永続ID', '氏名（参考）', 'グループ名（参考）']

CATEGORIES = ('admin', 'regular', 'guest')


# ────────────────────────────────────────────
# 小道具
# ────────────────────────────────────────────

def _now():
    return datetime.now(JST).replace(tzinfo=None)


def _s(v):
    return '' if v is None else str(v).strip()


def _is_admin():
    return check_is_total_admin(session.get('user_id'))


def _deny():
    return jsonify({'success': False, 'error': '権限がありません（永続IDの操作は総管理者のみ）'}), 403


def _cfg(name, default):
    """config.py の定数を読む（未定義のサイトでも動くように getattr で拾う）"""
    try:
        from config import Config
        return getattr(Config, name, default)
    except Exception:
        return default


def is_issuer():
    return bool(_cfg('PID_ISSUER', False))


def prefix():
    p = _s(_cfg('PID_PREFIX', '')).upper()
    return p if re.match(r'^[A-Z0-9]{2,8}$', p) else ''


def _site():
    try:
        return request.host
    except Exception:
        return ''


def _me_label():
    return f"{_s(session.get('full_name'))}({session.get('user_id')})".strip()


# ────────────────────────────────────────────
# 発番
# ────────────────────────────────────────────

def _next_seq(cursor, pre, kind):
    tag = KIND_TAG[kind]
    cursor.execute("SELECT pid FROM ug_pids WHERE pid LIKE %s", (f"{pre}-{tag}-%",))
    pat = re.compile(rf'^{re.escape(pre)}-{re.escape(tag)}-(\d+)$')
    mx = 0
    for r in cursor.fetchall():
        m = pat.match(r['pid'] or '')
        if m:
            mx = max(mx, int(m.group(1)))
    return mx + 1


def _mint(cursor, pre, kind, seq):
    return f"{pre}-{KIND_TAG[kind]}-{seq:06d}"


@user_groups_bp.route('/api/pids/issue', methods=['POST'])
@login_required
def pids_issue():
    """まだ永続IDの無い users / user_groups に発番する（源泉サイトのみ・冪等）"""
    if not _is_admin():
        return _deny()
    if not is_issuer():
        return jsonify({'success': False, 'error': 'このサイトは源泉サイトではありません（config.py の PID_ISSUER）'}), 400
    pre = prefix()
    if not pre:
        return jsonify({'success': False, 'error': 'config.py の PID_PREFIX が未設定です（英数2〜8文字）'}), 400

    conn = get_db()
    cursor = conn.cursor(dictionary=True)
    now = _now()
    me = _me_label()
    made = {'person': 0, 'group': 0}
    try:
        cursor.execute("""
            SELECT u.id, u.email, u.full_name FROM users u
            LEFT JOIN ug_pids p ON p.kind='person' AND p.local_id = u.id
            WHERE u.deleted_at IS NULL AND p.id IS NULL
            ORDER BY u.id
        """)
        targets = cursor.fetchall()          # 先に取り切る（同じカーソルを次の問い合わせで使うため）
        seq = _next_seq(cursor, pre, 'person')
        for u in targets:
            cursor.execute("""INSERT INTO ug_pids (kind, pid, local_id, email, display_name,
                                                   issued_at, issued_by, source_site)
                              VALUES ('person',%s,%s,%s,%s,%s,%s,%s)""",
                           (_mint(cursor, pre, 'person', seq), u['id'], (u['email'] or '').lower(),
                            u['full_name'], now, me, _site()))
            seq += 1
            made['person'] += 1

        cursor.execute("""
            SELECT g.id, g.name FROM user_groups g
            LEFT JOIN ug_pids p ON p.kind='group' AND p.local_id = g.id
            WHERE p.id IS NULL ORDER BY g.id
        """)
        targets = cursor.fetchall()
        seq = _next_seq(cursor, pre, 'group')
        for g in targets:
            cursor.execute("""INSERT INTO ug_pids (kind, pid, local_id, display_name,
                                                   issued_at, issued_by, source_site)
                              VALUES ('group',%s,%s,%s,%s,%s,%s)""",
                           (_mint(cursor, pre, 'group', seq), g['id'], g['name'], now, me, _site()))
            seq += 1
            made['group'] += 1
        conn.commit()
    except Exception as e:
        conn.rollback()
        logging.exception("pids_issue failed")
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        cursor.close()
        conn.close()
    return jsonify({'success': True, 'issued': made})


# ────────────────────────────────────────────
# 概要
# ────────────────────────────────────────────

def _counts(cursor):
    out = {}
    cursor.execute("SELECT COUNT(*) AS n FROM users WHERE deleted_at IS NULL")
    out['users'] = cursor.fetchone()['n']
    cursor.execute("""SELECT COUNT(*) AS n FROM ug_pids p JOIN users u ON u.id=p.local_id
                      WHERE p.kind='person' AND u.deleted_at IS NULL""")
    out['users_with_pid'] = cursor.fetchone()['n']
    cursor.execute("SELECT COUNT(*) AS n FROM user_groups")
    out['groups'] = cursor.fetchone()['n']
    cursor.execute("""SELECT COUNT(*) AS n FROM ug_pids p JOIN user_groups g ON g.id=p.local_id
                      WHERE p.kind='group'""")
    out['groups_with_pid'] = cursor.fetchone()['n']
    cursor.execute("SELECT COUNT(*) AS n FROM user_group_memberships")
    out['memberships'] = cursor.fetchone()['n']
    cursor.execute("SELECT COUNT(*) AS n FROM ug_pids WHERE local_id IS NULL")
    out['pending'] = cursor.fetchone()['n']
    return out


@user_groups_bp.route('/api/pids/summary')
@login_required
def pids_summary():
    if not _is_admin():
        return _deny()
    conn = get_db()
    cursor = conn.cursor(dictionary=True)
    try:
        c = _counts(cursor)
        try:
            cursor.execute("""SELECT version, source_site, imported_at, file_name,
                                     persons_linked, persons_created, groups_linked, groups_created,
                                     memberships_added, imported_by
                              FROM ug_pid_imports ORDER BY id DESC LIMIT 10""")
            imports = [{k: (v.strftime('%Y-%m-%d %H:%M') if isinstance(v, datetime) else v)
                        for k, v in r.items()} for r in cursor.fetchall()]
        except Exception:
            imports = []
    finally:
        cursor.close()
        conn.close()
    return jsonify({'success': True, 'counts': c, 'issuer': is_issuer(),
                    'prefix': prefix(), 'site': _site(), 'imports': imports})


# ────────────────────────────────────────────
# エクスポート（XLSX）
# ────────────────────────────────────────────

def _version():
    return f"{_now().strftime('%Y%m%d-%H%M%S')}"


def _export_rows(cursor):
    cursor.execute("""
        SELECT p.pid, u.email, u.full_name, u.category, u.affiliation
        FROM ug_pids p JOIN users u ON u.id = p.local_id
        WHERE p.kind='person' AND u.deleted_at IS NULL
        ORDER BY p.pid
    """)
    persons = [[r['pid'], (r['email'] or '').lower(), r['full_name'] or '',
                r['category'] or '', r['affiliation'] or '', ''] for r in cursor.fetchall()]

    cursor.execute("""
        SELECT p.pid, g.name, g.description
        FROM ug_pids p JOIN user_groups g ON g.id = p.local_id
        WHERE p.kind='group' ORDER BY p.pid
    """)
    groups = [[r['pid'], r['name'] or '', r['description'] or '', ''] for r in cursor.fetchall()]

    # 所属は「直接メンバー ∪ ルール由来 − 除外」を展開した結果を書き出す．
    # 台帳のグループは ug_group_rules から導かれるので，user_group_memberships を
    # そのまま読むと直接メンバーしか出ない．utils と同じ解釈を使う．
    from .utils import _group_member_ids
    at = _now()

    cursor.execute("""
        SELECT p.pid, p.local_id, u.full_name
        FROM ug_pids p JOIN users u ON u.id = p.local_id
        WHERE p.kind='person' AND u.deleted_at IS NULL
    """)
    person_pid, person_name = {}, {}
    for r in cursor.fetchall():
        person_pid[r['local_id']] = r['pid']
        person_name[r['local_id']] = r['full_name'] or ''

    cursor.execute("""
        SELECT p.pid, p.local_id, g.name
        FROM ug_pids p JOIN user_groups g ON g.id = p.local_id
        WHERE p.kind='group' ORDER BY p.pid
    """)
    group_rows = cursor.fetchall()

    members = []
    for g in group_rows:
        try:
            ids = _group_member_ids(cursor, g['local_id'], at)
        except Exception as e:
            logging.warning("group %s expand failed: %s", g['pid'], e)
            continue
        for uid in sorted(ids):
            if uid in person_pid:
                members.append([person_pid[uid], g['pid'], person_name[uid], g['name'] or ''])
    return persons, groups, members


def _workbook(version, persons, groups, members, counts):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_font = Font(bold=True)
    head_fill = PatternFill('solid', fgColor='DDEBF7')

    def sheet(title, headers, rows, widths):
        ws = wb.active if wb.active.title == 'Sheet' else wb.create_sheet()
        ws.title = title
        ws.append(headers)
        for c in ws[1]:
            c.font = head_font
            c.fill = head_fill
            c.alignment = Alignment(vertical='center')
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = 'A2'
        return ws

    info = [
        ['版', version],
        ['発行元サイト', _site()],
        ['発行日時', _now().strftime('%Y-%m-%d %H:%M:%S')],
        ['接頭辞', prefix()],
        ['源泉サイト', 'はい' if is_issuer() else 'いいえ'],
        ['人', len(persons)],
        ['グループ', len(groups)],
        ['所属', len(members)],
        ['永続ID未発番の人', counts['users'] - counts['users_with_pid']],
        ['永続ID未発番のグループ', counts['groups'] - counts['groups_with_pid']],
        ['注記', '取り込みは追加のみ．ここに載っていない所属を削除することはありません'],
        ['注記', 'ユーザ区分は受け入れ側を優先します．新規に作る人のみ admin は regular になります'],
    ]
    sheet(SHEET_INFO, INFO_HEADERS, info, [26, 60])
    sheet(SHEET_PERSON, PERSON_HEADERS, persons, [18, 36, 18, 10, 18, 30])
    sheet(SHEET_GROUP, GROUP_HEADERS, groups, [18, 32, 46, 30])
    sheet(SHEET_MEMBER, MEMBER_HEADERS, members, [18, 18, 18, 32])
    return wb


@user_groups_bp.route('/pids/export.xlsx')
@login_required
def pids_export():
    if not _is_admin():
        return _deny()
    conn = get_db()
    cursor = conn.cursor(dictionary=True)
    try:
        counts = _counts(cursor)
        persons, groups, members = _export_rows(cursor)
    finally:
        cursor.close()
        conn.close()
    version = _version()
    wb = _workbook(version, persons, groups, members, counts)
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return send_file(bio, as_attachment=True,
                     download_name=f"fujinp_pids_{version}.xlsx",
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


# ────────────────────────────────────────────
# インポート（XLSX）
# ────────────────────────────────────────────

def _read_sheet(ws, headers, key):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    head = [_s(h) for h in rows[0]]
    idx = {h: head.index(h) for h in headers if h in head}
    if key not in idx:
        raise ValueError(f"シート「{ws.title}」に「{key}」列がありません")
    out = []
    for n, r in enumerate(rows[1:], start=2):
        if r is None or all(_s(v) == '' for v in r):
            continue
        d = {h: (r[i] if i < len(r) else None) for h, i in idx.items()}
        d['_row'] = n
        out.append(d)
    return out


def parse_workbook(file_obj):
    from openpyxl import load_workbook
    wb = load_workbook(file_obj, data_only=True)
    warnings = []

    meta = {}
    if SHEET_INFO in wb.sheetnames:
        for r in wb[SHEET_INFO].iter_rows(values_only=True):
            if r and _s(r[0]) and len(r) > 1:
                meta.setdefault(_s(r[0]), _s(r[1]))

    persons, groups, members = [], [], []
    if SHEET_PERSON in wb.sheetnames:
        for d in _read_sheet(wb[SHEET_PERSON], PERSON_HEADERS, '永続ID'):
            pid = _s(d.get('永続ID'))
            if not PID_RE.match(pid):
                warnings.append(f"人 {d['_row']} 行目：永続ID「{pid}」の形式が不正なので飛ばします")
                continue
            cat = _s(d.get('区分')).lower()
            persons.append({
                'pid': pid,
                'email': _s(d.get('メール')).lower(),
                'full_name': _s(d.get('氏名')),
                'category': cat if cat in CATEGORIES else 'guest',
                'affiliation': _s(d.get('所属')),
                'row': d['_row'],
            })
    else:
        warnings.append('シート「人」がありません')

    if SHEET_GROUP in wb.sheetnames:
        for d in _read_sheet(wb[SHEET_GROUP], GROUP_HEADERS, '永続ID'):
            pid = _s(d.get('永続ID'))
            if not PID_RE.match(pid):
                warnings.append(f"グループ {d['_row']} 行目：永続ID「{pid}」の形式が不正なので飛ばします")
                continue
            groups.append({'pid': pid, 'name': _s(d.get('グループ名')),
                           'description': _s(d.get('説明')), 'row': d['_row']})

    if SHEET_MEMBER in wb.sheetnames:
        for d in _read_sheet(wb[SHEET_MEMBER], MEMBER_HEADERS, '人の永続ID'):
            ppid, gpid = _s(d.get('人の永続ID')), _s(d.get('グループの永続ID'))
            if not (PID_RE.match(ppid) and PID_RE.match(gpid)):
                warnings.append(f"所属 {d['_row']} 行目：永続IDの形式が不正なので飛ばします")
                continue
            members.append({'person_pid': ppid, 'group_pid': gpid, 'row': d['_row']})

    return {'meta': meta, 'persons': persons, 'groups': groups,
            'members': members, 'warnings': warnings}


def _pid_map(cursor, kind):
    cursor.execute("SELECT pid, local_id FROM ug_pids WHERE kind=%s", (kind,))
    return {r['pid']: r['local_id'] for r in cursor.fetchall()}


@user_groups_bp.route('/api/pids/import', methods=['POST'])
@login_required
def pids_import():
    """
    multipart: file=xlsx, mode=check|apply
      check … DB に触れず，何が起きるかを数えて返す
      apply … 追加のみ．既存の users の区分・氏名・所属は書き換えない．所属の削除もしない
    """
    if not _is_admin():
        return _deny()
    f = request.files.get('file')
    mode = request.form.get('mode', 'check')
    if not f:
        return jsonify({'success': False, 'error': 'ファイルがありません'}), 400
    fname = f.filename or ''
    try:
        parsed = parse_workbook(f.stream)
    except Exception as e:
        return jsonify({'success': False, 'error': f'読み取り失敗: {e}'}), 400

    conn = get_db()
    cursor = conn.cursor(dictionary=True)
    now = _now()
    me_id = session.get('user_id')
    apply_ = (mode == 'apply')
    st = {'persons_known': 0, 'persons_linked': 0, 'persons_created': 0,
          'persons_pending': 0, 'persons_demoted': 0,
          'groups_known': 0, 'groups_linked': 0, 'groups_created': 0,
          'memberships_added': 0, 'memberships_exists': 0, 'memberships_skipped': 0}
    notes = list(parsed['warnings'])
    try:
        pmap = _pid_map(cursor, 'person')
        gmap = _pid_map(cursor, 'group')

        # ── 人 ──
        for p in parsed['persons']:
            local = pmap.get(p['pid'])
            if local:
                st['persons_known'] += 1
                continue
            u = None
            if p['email']:
                cursor.execute("SELECT id, category FROM users WHERE email=%s AND deleted_at IS NULL", (p['email'],))
                u = cursor.fetchone()
            if u:
                st['persons_linked'] += 1
                if apply_:
                    _upsert_pid(cursor, 'person', p['pid'], u['id'], p['email'], p['full_name'], now, parsed['meta'])
                pmap[p['pid']] = u['id']
                continue
            if not p['email']:
                st['persons_pending'] += 1
                notes.append(f"{p['pid']}：メールが空なのでユーザは作らず，永続IDだけ控えます")
                if apply_:
                    _upsert_pid(cursor, 'person', p['pid'], None, '', p['full_name'], now, parsed['meta'])
                continue
            cat = 'regular' if p['category'] == 'admin' else p['category']
            if cat != p['category']:
                st['persons_demoted'] += 1
            st['persons_created'] += 1
            if apply_:
                cursor.execute("""INSERT INTO users (email, full_name, category, affiliation,
                                                     is_active, created_at, updated_at)
                                  VALUES (%s,%s,%s,%s,TRUE,%s,%s)""",
                               (p['email'], p['full_name'], cat, p['affiliation'] or None, now, now))
                uid = cursor.lastrowid
                _upsert_pid(cursor, 'person', p['pid'], uid, p['email'], p['full_name'], now, parsed['meta'])
                pmap[p['pid']] = uid
            else:
                pmap[p['pid']] = -1   # 確認モード：これから作る人（所属の件数を数えるための仮の印）

        # ── グループ ──
        for g in parsed['groups']:
            if gmap.get(g['pid']):
                st['groups_known'] += 1
                continue
            cursor.execute("SELECT id FROM user_groups WHERE name=%s ORDER BY id LIMIT 1", (g['name'],))
            row = cursor.fetchone()
            if row:
                st['groups_linked'] += 1
                if apply_:
                    _upsert_pid(cursor, 'group', g['pid'], row['id'], '', g['name'], now, parsed['meta'])
                gmap[g['pid']] = row['id']
                continue
            st['groups_created'] += 1
            if apply_:
                cursor.execute("""INSERT INTO user_groups (name, description, manager_user_id, created_at, updated_at)
                                  VALUES (%s,%s,%s,%s,%s)""",
                               (g['name'], g['description'] or None, me_id, now, now))
                gid = cursor.lastrowid
                _upsert_pid(cursor, 'group', g['pid'], gid, '', g['name'], now, parsed['meta'])
                gmap[g['pid']] = gid
            else:
                gmap[g['pid']] = -1   # 確認モード：これから作るグループ

        # ── 所属（追加のみ） ──
        for m in parsed['members']:
            uid, gid = pmap.get(m['person_pid']), gmap.get(m['group_pid'])
            if not uid or not gid:
                st['memberships_skipped'] += 1
                continue
            if uid > 0 and gid > 0:
                cursor.execute("SELECT id FROM user_group_memberships WHERE group_id=%s AND user_id=%s LIMIT 1",
                               (gid, uid))
                exists = cursor.fetchone()
            else:
                exists = None       # まだ無いものなので所属も無い
            if exists:
                st['memberships_exists'] += 1
                continue
            st['memberships_added'] += 1
            if apply_:
                cursor.execute("""INSERT INTO user_group_memberships (group_id, user_id, created_at)
                                  VALUES (%s,%s,%s)""", (gid, uid, now))

        if apply_:
            try:
                cursor.execute("""INSERT INTO ug_pid_imports
                                  (version, source_site, file_name, persons_linked, persons_created,
                                   groups_linked, groups_created, memberships_added, imported_by, imported_at)
                                  VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)""",
                               (parsed['meta'].get('版', ''), parsed['meta'].get('発行元サイト', ''), fname,
                                st['persons_linked'], st['persons_created'],
                                st['groups_linked'], st['groups_created'],
                                st['memberships_added'], _me_label(), now))
            except Exception as e:
                logging.warning("ug_pid_imports unavailable: %s", e)
            conn.commit()
        else:
            conn.rollback()
    except Exception as e:
        conn.rollback()
        logging.exception("pids_import failed")
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        cursor.close()
        conn.close()

    return jsonify({'success': True, 'mode': 'apply' if apply_ else 'check',
                    'meta': parsed['meta'], 'summary': st, 'warnings': notes[:40]})


def _upsert_pid(cursor, kind, pid, local_id, email, name, now, meta):
    """永続IDの対応を控える．同じ永続IDの行があれば local_id を埋める"""
    cursor.execute("SELECT id FROM ug_pids WHERE pid=%s", (pid,))
    row = cursor.fetchone()
    if row:
        cursor.execute("""UPDATE ug_pids SET local_id=%s, email=COALESCE(NULLIF(%s,''), email),
                          display_name=COALESCE(NULLIF(%s,''), display_name) WHERE id=%s""",
                       (local_id, email, name, row['id']))
        return
    cursor.execute("""INSERT INTO ug_pids (kind, pid, local_id, email, display_name,
                                           issued_at, issued_by, source_site, note)
                      VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s)""",
                   (kind, pid, local_id, email or None, name or None, now,
                    '取り込み', meta.get('発行元サイト', ''), '配布物から取り込み'))


# ────────────────────────────────────────────
# 画面
# ────────────────────────────────────────────

PAGE = """
<!DOCTYPE html>
<html lang="ja"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>永続ID — FUJIN-P</title>
<style>
 body{font-family:-apple-system,"Hiragino Kaku Gothic ProN","Yu Gothic",Meiryo,sans-serif;
      margin:0;background:#f5f6f8;color:#222}
 header{background:#54367C;color:#fff;padding:14px 20px;display:flex;align-items:center;gap:14px}
 header h1{font-size:18px;margin:0;font-weight:600}
 header .sp{flex:1}
 header a{color:#fff;text-decoration:none;font-size:13px;border:1px solid rgba(255,255,255,.6);
          padding:5px 10px;border-radius:6px}
 main{max-width:980px;margin:20px auto;padding:0 16px}
 .card{background:#fff;border-radius:10px;padding:18px 20px;margin-bottom:16px;
       box-shadow:0 1px 3px rgba(0,0,0,.08)}
 .card h2{font-size:15px;margin:0 0 10px;color:#54367C}
 .muted{color:#666;font-size:13px;line-height:1.7}
 .grid{display:flex;flex-wrap:wrap;gap:10px;margin:10px 0}
 .stat{background:#f2f0f7;border-radius:8px;padding:10px 14px;min-width:120px}
 .stat b{display:block;font-size:20px}
 .stat span{font-size:12px;color:#555}
 button,.btn{background:#54367C;color:#fff;border:0;border-radius:6px;padding:8px 14px;
             font-size:14px;cursor:pointer;text-decoration:none;display:inline-block}
 button.sub{background:#fff;color:#54367C;border:1px solid #54367C}
 button:disabled{background:#bbb;cursor:not-allowed}
 table{border-collapse:collapse;width:100%;font-size:13px;margin-top:8px}
 th,td{border:1px solid #e3e3e8;padding:6px 8px;text-align:left}
 th{background:#f2f0f7}
 pre{background:#f7f7fa;border:1px solid #e3e3e8;border-radius:6px;padding:10px;
     font-size:12px;white-space:pre-wrap;max-height:320px;overflow:auto}
 .warn{color:#a4501a}
 .tag{display:inline-block;background:#e7e2f1;color:#54367C;border-radius:4px;
      padding:2px 8px;font-size:12px;margin-left:6px}
</style></head><body>
<header>
  <h1>永続ID</h1><span id="mode" class="tag"></span>
  <div class="sp"></div>
  <a href="{{ back }}">FUJINダッシュボードに戻る</a>
</header>
<main>
  <div class="card">
    <h2>このサイトの状況</h2>
    <div class="grid" id="stats"></div>
    <div class="muted" id="site"></div>
  </div>

  <div class="card" id="issue-card" style="display:none">
    <h2>発番</h2>
    <div class="muted">まだ永続IDの無いユーザとグループに番号を振ります．何度押しても同じ結果になります（すでに番号のあるものは飛ばします）．</div>
    <div class="grid"><button id="btn-issue">永続IDを発番する</button></div>
    <pre id="issue-out" style="display:none"></pre>
  </div>

  <div class="card">
    <h2>配布物を書き出す</h2>
    <div class="muted">永続ID・人・グループ・所属を XLSX で書き出します．受け取った側はこれを取り込むだけです．</div>
    <div class="grid"><a class="btn" href="{{ url_export }}">XLSX をダウンロード</a></div>
  </div>

  <div class="card">
    <h2>配布物を取り込む</h2>
    <div class="muted">
      追加のみです．ここに載っていない所属を消すことはありません．<br>
      すでにあるユーザの区分・氏名・所属は書き換えません．新しく作る人だけ，admin は regular にします．
    </div>
    <div class="grid">
      <input type="file" id="file" accept=".xlsx">
      <button id="btn-check" class="sub">確認する</button>
      <button id="btn-apply" disabled>取り込む</button>
    </div>
    <pre id="imp-out" style="display:none"></pre>
  </div>

  <div class="card">
    <h2>取り込んだ版</h2>
    <div id="imports" class="muted">—</div>
  </div>
</main>
<script>
const $ = s => document.querySelector(s);

async function load(){
  const r = await fetch('{{ url_summary }}');
  const j = await r.json();
  if(!j.success){ $('#stats').textContent = j.error; return; }
  const c = j.counts;
  $('#mode').textContent = j.issuer ? '源泉サイト（発番あり）' : '受け取り側';
  $('#site').textContent = 'サイト: ' + j.site + '　接頭辞: ' + (j.prefix || '（未設定）');
  $('#stats').innerHTML =
    stat(c.users_with_pid + ' / ' + c.users, '永続IDのある人') +
    stat(c.groups_with_pid + ' / ' + c.groups, '永続IDのあるグループ') +
    stat(c.memberships, '所属') +
    stat(c.pending, 'ローカル未対応の永続ID');
  if(j.issuer) $('#issue-card').style.display = '';
  const im = j.imports || [];
  $('#imports').innerHTML = im.length ? (
    '<table><tr><th>版</th><th>発行元</th><th>日時</th><th>人（紐づけ/作成）</th>'
    + '<th>グループ（紐づけ/作成）</th><th>所属追加</th></tr>' +
    im.map(r => '<tr><td>' + esc(r.version) + '</td><td>' + esc(r.source_site) + '</td><td>'
      + esc(r.imported_at) + '</td><td>' + r.persons_linked + ' / ' + r.persons_created
      + '</td><td>' + r.groups_linked + ' / ' + r.groups_created + '</td><td>'
      + r.memberships_added + '</td></tr>').join('') + '</table>') : 'まだありません';
}
const stat = (n,l) => '<div class="stat"><b>'+n+'</b><span>'+l+'</span></div>';
const esc = s => (s==null?'':String(s)).replace(/[&<>]/g, m => ({'&':'&amp;','<':'&lt;','>':'&gt;'}[m]));

$('#btn-issue').onclick = async () => {
  if(!confirm('まだ永続IDの無いユーザとグループに発番します．よろしいですか？')) return;
  $('#btn-issue').disabled = true;
  const r = await fetch('{{ url_issue }}', {method:'POST'});
  const j = await r.json();
  $('#issue-out').style.display = '';
  $('#issue-out').textContent = j.success
    ? ('発番しました　人 ' + j.issued.person + ' 件／グループ ' + j.issued.group + ' 件')
    : ('失敗: ' + j.error);
  $('#btn-issue').disabled = false;
  load();
};

async function send(mode){
  const f = $('#file').files[0];
  if(!f){ alert('ファイルを選んでください'); return; }
  const fd = new FormData();
  fd.append('file', f); fd.append('mode', mode);
  const r = await fetch('{{ url_import }}', {method:'POST', body:fd});
  const j = await r.json();
  $('#imp-out').style.display = '';
  if(!j.success){ $('#imp-out').textContent = '失敗: ' + j.error; return; }
  const s = j.summary;
  const lines = [
    (mode === 'apply' ? '── 取り込みました ──' : '── 確認（まだ書き込んでいません） ──'),
    '版: ' + (j.meta['版'] || '(不明)') + '　発行元: ' + (j.meta['発行元サイト'] || '(不明)'),
    '',
    '人　　　 既知 ' + s.persons_known + '／メールで紐づけ ' + s.persons_linked
      + '／新規作成 ' + s.persons_created + '（うち admin→regular ' + s.persons_demoted + '）'
      + '／保留 ' + s.persons_pending,
    'グループ 既知 ' + s.groups_known + '／名前で紐づけ ' + s.groups_linked
      + '／新規作成 ' + s.groups_created,
    '所属　　 追加 ' + s.memberships_added + '／すでにある ' + s.memberships_exists
      + '／永続IDが解決できず飛ばした ' + s.memberships_skipped,
  ];
  if(j.warnings && j.warnings.length) lines.push('', '注意:', ...j.warnings.map(w => '  ' + w));
  $('#imp-out').textContent = lines.join('\\n');
  $('#btn-apply').disabled = (mode === 'apply');
  if(mode === 'apply') load();
}
$('#btn-check').onclick = () => send('check');
$('#btn-apply').onclick = () => { if(confirm('取り込みます．追加のみで，既存のものは消しません．')) send('apply'); };
$('#file').onchange = () => { $('#btn-apply').disabled = false; };
load();
</script></body></html>
"""


@user_groups_bp.route('/pids')
@login_required
def pids_page():
    if not _is_admin():
        return _deny()
    from flask import url_for
    return render_template_string(
        PAGE,
        back=url_for('user_groups.return_to_fujin'),
        url_summary=url_for('user_groups.pids_summary'),
        url_issue=url_for('user_groups.pids_issue'),
        url_import=url_for('user_groups.pids_import'),
        url_export=url_for('user_groups.pids_export'),
    )
